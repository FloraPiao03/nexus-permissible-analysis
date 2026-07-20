#!/usr/bin/env python3
"""Build an Interventional-only view from a validated five-category trials.csv."""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BUCKETS = ("UNKNOWN", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "NEXUS")
US_BUCKETS = ("US_ONLY", "US_NON_CHINA_MULTI", "NEXUS")
REQUIRED_FIELDS = (
    "nct_id", "study_type", "countries_pipe", "country_count", "has_location_data",
    "has_us", "has_china", "bucket", "us_presence_group",
)
SCOPE_VALUE = "INTERVENTIONAL"
COMPARISON_SHEET = "All_vs_Interventional_Compare"  # Excel worksheet names are limited to 31 characters.
LOCATION_LIMITATION = (
    "Registered or planned study facilities, not confirmed participant nationality "
    "or actual country-level enrollment."
)
BODY_FONT = Font(name="Arial", size=10)


def pct(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def read_source(path: Path) -> tuple[list[str], list[dict], Counter]:
    vocabulary = Counter()
    selected = []
    seen = set()
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        fields = reader.fieldnames or []
        missing = sorted(set(REQUIRED_FIELDS) - set(fields))
        if missing:
            raise ValueError(f"Source trials.csv is missing required fields: {missing}")
        for row in reader:
            nct_id = row["nct_id"]
            if not nct_id or nct_id in seen:
                raise ValueError(f"Missing or duplicate NCT ID in validated source: {nct_id!r}")
            seen.add(nct_id)
            vocabulary[row["study_type"]] += 1
            if row["study_type"] == SCOPE_VALUE:
                selected.append(row)
    return fields, selected, vocabulary


def scan_source(path: Path) -> tuple[list[str], dict, Counter]:
    """Aggregate the selected scope without retaining hundreds of thousands of rows."""
    vocabulary = Counter()
    counts = Counter()
    groups = Counter()
    seen = set()
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        fields = reader.fieldnames or []
        missing = sorted(set(REQUIRED_FIELDS) - set(fields))
        if missing:
            raise ValueError(f"Source trials.csv is missing required fields: {missing}")
        for row in reader:
            nct_id = row["nct_id"]
            if not nct_id or nct_id in seen:
                raise ValueError(f"Missing or duplicate NCT ID in validated source: {nct_id!r}")
            seen.add(nct_id)
            vocabulary[row["study_type"]] += 1
            if row["study_type"] == SCOPE_VALUE:
                counts[row["bucket"]] += 1
                groups[row["us_presence_group"]] += 1
    summary = aggregate_counts(counts, groups)
    if any(summary["us_presence_groups"][group]["count"] != groups[group] for group in ("UNKNOWN", "NO_US", "HAS_US")):
        raise AssertionError("Source us_presence_group counts disagree with bucket-derived groups")
    return fields, summary, vocabulary


def iter_selected(path: Path):
    with path.open(newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["study_type"] == SCOPE_VALUE:
                yield row


def aggregate(rows: list[dict]) -> dict:
    counts = Counter(row["bucket"] for row in rows)
    groups = Counter(row["us_presence_group"] for row in rows)
    return aggregate_counts(counts, groups)


def aggregate_counts(counts: Counter, groups: Counter) -> dict:
    unexpected_buckets = sorted(set(counts) - set(BUCKETS))
    unexpected_groups = sorted(set(groups) - {"UNKNOWN", "NO_US", "HAS_US"})
    if unexpected_buckets or unexpected_groups:
        raise AssertionError(f"Unexpected classifications: {unexpected_buckets}, {unexpected_groups}")
    total = sum(counts.values())
    unknown = counts["UNKNOWN"]
    known = total - unknown
    has_us = sum(counts[b] for b in US_BUCKETS)
    checks = {
        "five_buckets_equal_total": sum(counts[b] for b in BUCKETS) - total,
        "permissible_plus_has_us_equal_known": counts["PERMISSIBLE"] + has_us - known,
        "us_subcategories_equal_has_us": sum(counts[b] for b in US_BUCKETS) - has_us,
        "row_group_counts_match_bucket_groups": {
            "UNKNOWN": groups["UNKNOWN"] - unknown,
            "NO_US": groups["NO_US"] - counts["PERMISSIBLE"],
            "HAS_US": groups["HAS_US"] - has_us,
        },
    }
    if any(value != 0 for key, value in checks.items() if key != "row_group_counts_match_bucket_groups"):
        raise AssertionError(f"Interventional reconciliation failed: {checks}")
    if any(checks["row_group_counts_match_bucket_groups"].values()):
        raise AssertionError(f"Interventional group mapping failed: {checks}")
    buckets = {}
    for bucket in BUCKETS:
        count = counts[bucket]
        buckets[bucket] = {
            "count": count,
            "percentage_total_interventional": pct(count, total),
            "percentage_known_location_interventional": None if bucket == "UNKNOWN" else pct(count, known),
            "percentage_interventional_has_us": pct(count, has_us) if bucket in US_BUCKETS else None,
        }
    return {
        "total_studies": total,
        "known_location_studies": known,
        "unknown_location_studies": unknown,
        "buckets": buckets,
        "us_presence_groups": {
            "UNKNOWN": {"count": unknown, "percentage_total_interventional": pct(unknown, total)},
            "NO_US": {"count": counts["PERMISSIBLE"], "percentage_total_interventional": pct(counts["PERMISSIBLE"], total), "percentage_known_location_interventional": pct(counts["PERMISSIBLE"], known)},
            "HAS_US": {"count": has_us, "percentage_total_interventional": pct(has_us, total), "percentage_known_location_interventional": pct(has_us, known)},
        },
        "reconciliation_differences": checks,
    }


def load_all_summary(path: Path) -> dict:
    summary = json.loads(path.read_text(encoding="utf-8"))
    if summary.get("classification_model_version") != 2:
        raise ValueError("All-study summary is not the validated five-category model")
    return summary


def comparison_rows(all_summary: dict, interventional: dict) -> list[dict]:
    all_total = all_summary["total_studies"]
    int_total = interventional["total_studies"]
    rows = [
        {"metric": "Total denominator", "all_count": all_total, "all_pct": 1.0, "int_count": int_total, "int_pct": 1.0},
        {"metric": "Known-location denominator", "all_count": all_summary["known_location_studies"], "all_pct": pct(all_summary["known_location_studies"], all_total), "int_count": interventional["known_location_studies"], "int_pct": pct(interventional["known_location_studies"], int_total)},
    ]
    for bucket in BUCKETS:
        rows.append({
            "metric": bucket,
            "all_count": all_summary["buckets"][bucket]["count"],
            "all_pct": all_summary["buckets"][bucket]["percentage_total"],
            "int_count": interventional["buckets"][bucket]["count"],
            "int_pct": interventional["buckets"][bucket]["percentage_total_interventional"],
        })
    rows.append({
        "metric": "HAS_US",
        "all_count": all_summary["us_presence_groups"]["HAS_US"]["count"],
        "all_pct": all_summary["us_presence_groups"]["HAS_US"]["percentage_total"],
        "int_count": interventional["us_presence_groups"]["HAS_US"]["count"],
        "int_pct": interventional["us_presence_groups"]["HAS_US"]["percentage_total_interventional"],
    })
    for row in rows:
        row["count_difference_int_minus_all"] = row["int_count"] - row["all_count"]
        row["percentage_point_difference"] = row["int_pct"] - row["all_pct"]
        row["percentage_denominator"] = "Each scope's total denominator"
    return rows


def write_csv(path: Path, fields: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in fields} for row in rows)


def prepare_sheet(ws, headers: list[str], widths: list[int]) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    cells = []
    for value in headers:
        cell = WriteOnlyCell(ws, value=value)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)
        cells.append(cell)
    ws.append(cells)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def append_row(ws, values, percentage_columns=(), style_body=True) -> None:
    cells = []
    for index, value in enumerate(values, 1):
        cell = WriteOnlyCell(ws, value=value)
        if style_body:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 60)
        if index in percentage_columns and value is not None:
            cell.number_format = "0.00%"
        cells.append(cell)
    ws.append(cells)


def make_workbook(path: Path, summary: dict, detail_csv: Path, fields: list[str], comparison: list[dict], metadata: dict) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Executive_Summary")
    prepare_sheet(ws, ["Interventional-only Nexus & Permissible Analysis", "Value"], [58, 100])
    items = [
        ("Analysis scope", 'ClinicalTrials.gov studies with study_type == "INTERVENTIONAL"'),
        ("Primary denominator", "All unique INTERVENTIONAL studies in the validated frozen snapshot"),
        ("Location interpretation", LOCATION_LIMITATION),
        ("Total Interventional studies", summary["total_studies"]),
        ("Known-location Interventional studies", summary["known_location_studies"]),
        ("UNKNOWN Interventional studies", summary["unknown_location_studies"]),
        ("PERMISSIBLE count", summary["buckets"]["PERMISSIBLE"]["count"]),
        ("PERMISSIBLE % of all Interventional", summary["buckets"]["PERMISSIBLE"]["percentage_total_interventional"]),
        ("PERMISSIBLE % of known-location Interventional", summary["buckets"]["PERMISSIBLE"]["percentage_known_location_interventional"]),
        ("HAS_US count", summary["us_presence_groups"]["HAS_US"]["count"]),
        ("HAS_US % of all Interventional", summary["us_presence_groups"]["HAS_US"]["percentage_total_interventional"]),
        ("NEXUS count", summary["buckets"]["NEXUS"]["count"]),
        ("NEXUS % of all Interventional", summary["buckets"]["NEXUS"]["percentage_total_interventional"]),
        ("NEXUS % of known-location Interventional", summary["buckets"]["NEXUS"]["percentage_known_location_interventional"]),
        ("NEXUS % of Interventional HAS_US", summary["buckets"]["NEXUS"]["percentage_interventional_has_us"]),
    ]
    for label, value in items:
        append_row(ws, [label, value], (2,) if isinstance(value, float) else ())

    cs = wb.create_sheet("Classification_Summary")
    prepare_sheet(cs, ["bucket", "count", "% all Interventional", "% known-location Interventional", "% Interventional HAS_US", "denominator notes"], [25, 16, 23, 31, 28, 55])
    for bucket in BUCKETS:
        item = summary["buckets"][bucket]
        append_row(cs, [bucket, item["count"], item["percentage_total_interventional"], item["percentage_known_location_interventional"], item["percentage_interventional_has_us"], "Percentages use the explicitly named Interventional denominator"], (3, 4, 5))

    detail = wb.create_sheet("Study_Detail")
    prepare_sheet(detail, fields, [16, 45, 18, 20, 45, 14, 18, 12, 14, 24, 20, 18, 14, 14, 20][:len(fields)])
    with detail_csv.open(newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            append_row(detail, [row.get(field, "") for field in fields], style_body=False)

    defs = wb.create_sheet("Definitions")
    prepare_sheet(defs, ["Item", "Definition"], [30, 105])
    definitions = [
        ("Analysis scope", 'ClinicalTrials.gov studies with study_type == "INTERVENTIONAL"'),
        ("UNKNOWN", "No usable registered location-country value"),
        ("PERMISSIBLE / NO_US", "At least one usable country and no United States"),
        ("US_ONLY", "Complete unique country set is exactly {United States}"),
        ("US_NON_CHINA_MULTI", "Contains United States, not China, and at least one additional non-US country"),
        ("NEXUS", "Contains both United States and China; other countries may also be present"),
        ("HAS_US", "US_ONLY + US_NON_CHINA_MULTI + NEXUS"),
        ("Location interpretation", LOCATION_LIMITATION),
    ]
    for item in definitions:
        append_row(defs, item)

    meta = wb.create_sheet("Run_Metadata")
    prepare_sheet(meta, ["Key", "Value"], [38, 105])
    for key, value in metadata.items():
        append_row(meta, [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])

    comp = wb.create_sheet(COMPARISON_SHEET)
    prepare_sheet(comp, ["Metric", "All Studies count", "All Studies %", "Interventional count", "Interventional %", "Count difference (Int - All)", "Percentage-point difference", "Percentage denominator"], [30, 20, 18, 22, 20, 28, 28, 34])
    for row in comparison:
        append_row(comp, [row["metric"], row["all_count"], row["all_pct"], row["int_count"], row["int_pct"], row["count_difference_int_minus_all"], row["percentage_point_difference"], row["percentage_denominator"]], (3, 5, 7))
    wb.save(path)


def verify_workbook(path: Path, summary: dict) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    expected = ["Executive_Summary", "Classification_Summary", "Study_Detail", "Definitions", "Run_Metadata", COMPARISON_SHEET]
    if wb.sheetnames != expected:
        raise AssertionError(f"Unexpected workbook sheets: {wb.sheetnames}")
    executive = {row[0].value: row[1].value for row in wb["Executive_Summary"].iter_rows(min_row=2) if len(row) >= 2}
    if executive.get("Total Interventional studies") != summary["total_studies"]:
        raise AssertionError("Excel total does not match Interventional summary")
    classification = {row[0].value: row[1].value for row in wb["Classification_Summary"].iter_rows(min_row=2)}
    for bucket in BUCKETS:
        if classification.get(bucket) != summary["buckets"][bucket]["count"]:
            raise AssertionError(f"Excel {bucket} count mismatch")
    detail_count = sum(1 for _ in wb["Study_Detail"].iter_rows(min_row=2, values_only=True))
    if detail_count != summary["total_studies"]:
        raise AssertionError("Excel Study_Detail row count mismatch")
    wb.close()


def markdown(summary: dict, comparison: list[dict]) -> str:
    f = lambda value: "N/A" if value is None else f"{value:.2%}"
    lines = [
        "# Interventional-only Nexus and Permissible Analysis", "",
        '> Analysis scope: ClinicalTrials.gov studies with `study_type == "INTERVENTIONAL"`', "",
        f"- Total Interventional studies: **{summary['total_studies']:,}**",
        f"- Known-location Interventional studies: **{summary['known_location_studies']:,}**",
        f"- UNKNOWN: **{summary['unknown_location_studies']:,} ({f(summary['buckets']['UNKNOWN']['percentage_total_interventional'])} of all Interventional)**", "",
        "## Five-category results", "",
        "| Category | Count | % all Interventional | % known-location Interventional | % Interventional HAS_US |",
        "|---|---:|---:|---:|---:|",
    ]
    for bucket in BUCKETS:
        item = summary["buckets"][bucket]
        lines.append(f"| {bucket} | {item['count']:,} | {f(item['percentage_total_interventional'])} | {f(item['percentage_known_location_interventional'])} | {f(item['percentage_interventional_has_us'])} |")
    lines += ["", "## High-level grouping", "",
              f"- NO_US / PERMISSIBLE: **{summary['us_presence_groups']['NO_US']['count']:,}**",
              f"- HAS_US: **{summary['us_presence_groups']['HAS_US']['count']:,} ({f(summary['us_presence_groups']['HAS_US']['percentage_total_interventional'])} of all Interventional)**", "",
              "## All studies vs Interventional only", "",
              "All percentages below use each scope's own total denominator.", "",
              "| Metric | All Studies | Interventional Only | Difference (Int - All) |",
              "|---|---:|---:|---:|"]
    for row in comparison:
        lines.append(f"| {row['metric']} | {row['all_count']:,} ({f(row['all_pct'])}) | {row['int_count']:,} ({f(row['int_pct'])}) | {row['count_difference_int_minus_all']:,}; {row['percentage_point_difference']:+.2%} pp |")
    checks = summary["reconciliation_differences"]
    lines += ["", "## Reconciliation checks", "",
              f"- Five buckets minus total: **{checks['five_buckets_equal_total']}**",
              f"- PERMISSIBLE + HAS_US minus known-location: **{checks['permissible_plus_has_us_equal_known']}**",
              f"- US subcategories minus HAS_US: **{checks['us_subcategories_equal_has_us']}**", "",
              "## Limitation", "", f"- Location interpretation: {LOCATION_LIMITATION}"]
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True, help="Validated revised out_v2/trials.csv")
    parser.add_argument("--all-summary", type=Path, required=True, help="Validated revised out_v2/summary.json")
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true", help="Replace this command's existing generated artifacts")
    parser.add_argument("--workbook-only", action="store_true", help="Rebuild Excel from existing generated CSV/JSON")
    args = parser.parse_args(argv)
    if args.workbook_only:
        summary = json.loads((args.outdir / "summary_interventional.json").read_text(encoding="utf-8"))
        detail_csv = args.outdir / "trials_interventional.csv"
        with detail_csv.open(newline="", encoding="utf-8") as fh:
            fields = csv.DictReader(fh).fieldnames or []
        metadata = {
            "analysis_scope": summary["analysis_scope"],
            "source_study_level_file": summary["source_study_level_file"],
            "source_all_study_summary": summary["source_all_study_summary"],
            "generated_at_utc": summary["generated_at_utc"],
            "classification_model_version": 2,
            "study_type_vocabulary_all_studies": summary["study_type_vocabulary_all_studies"],
            "raw_snapshot_modified": False,
            "api_reharvest_performed": False,
        }
        workbook = args.outdir / "nexus_permissible_interventional.xlsx"
        make_workbook(workbook, summary, detail_csv, fields, summary["all_vs_interventional_comparison"], metadata)
        verify_workbook(workbook, summary)
        print(workbook)
        return 0
    final_targets = [args.outdir / "nexus_permissible_interventional.xlsx", args.outdir / "interventional_audit.json"]
    if any(path.exists() for path in final_targets) and not args.overwrite:
        parser.error(f"completed output artifacts already exist in: {args.outdir}")
    fields, summary, vocabulary = scan_source(args.source)
    all_summary = load_all_summary(args.all_summary)
    comparison = comparison_rows(all_summary, summary)
    summary.update({
        "analysis_scope": 'study_type == "INTERVENTIONAL"',
        "source_study_level_file": str(args.source.resolve()),
        "source_all_study_summary": str(args.all_summary.resolve()),
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "classification_model_version": 2,
        "study_type_vocabulary_all_studies": dict(sorted(vocabulary.items())),
        "all_vs_interventional_comparison": comparison,
        "location_interpretation": LOCATION_LIMITATION,
    })
    args.outdir.mkdir(parents=True, exist_ok=True)
    detail_csv = args.outdir / "trials_interventional.csv"
    write_csv(detail_csv, fields, iter_selected(args.source))
    (args.outdir / "summary_interventional.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (args.outdir / "summary_interventional.md").write_text(markdown(summary, comparison), encoding="utf-8")
    workbook = args.outdir / "nexus_permissible_interventional.xlsx"
    metadata = {
        "analysis_scope": summary["analysis_scope"],
        "source_study_level_file": summary["source_study_level_file"],
        "source_all_study_summary": summary["source_all_study_summary"],
        "generated_at_utc": summary["generated_at_utc"],
        "classification_model_version": 2,
        "study_type_vocabulary_all_studies": summary["study_type_vocabulary_all_studies"],
        "raw_snapshot_modified": False,
        "api_reharvest_performed": False,
    }
    make_workbook(workbook, summary, detail_csv, fields, comparison, metadata)
    verify_workbook(workbook, summary)
    print(args.outdir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
