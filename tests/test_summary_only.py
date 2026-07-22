import hashlib
import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr
from pathlib import Path

from openpyxl import load_workbook

from analyze import main


def study(nct_id, start_date, countries, start_type="ACTUAL", sponsor_class="INDUSTRY", intervention_types=None):
    status = {"overallStatus": "COMPLETED"}
    if start_date is not None:
        status["startDateStruct"] = {"date": start_date, "type": start_type}
    return {"protocolSection": {
        "identificationModule": {"nctId": nct_id, "briefTitle": nct_id},
        "designModule": {"studyType": "INTERVENTIONAL"},
        "statusModule": status,
        "sponsorCollaboratorsModule": {"leadSponsor": {"class": sponsor_class}},
        "armsInterventionsModule": {"interventions": [{"type": value} for value in (intervention_types or [])]},
        "contactsLocationsModule": {"locations": [{"country": country} for country in countries]},
    }}


def make_run(root: Path, complete=True) -> Path:
    run = root / "run_fixture"
    raw = run / "raw"
    raw.mkdir(parents=True)
    pages = [
        [
            study("NCT00000001", "2025-01-01", ["United States", "China"], intervention_types=["DRUG"]),
            study("NCT00000002", "2024-02", ["Germany"], intervention_types=["DEVICE", "DRUG", "DRUG"]),
            study("NCT00000003", "2022", ["United States"], intervention_types=["OTHER", "DRUG"]),
            study("NCT00000004", "2021-01-01", ["United States", "Canada"], intervention_types=["DEVICE"]),
        ],
        [
            study("NCT00000005", "2019-01-01", []),
            study("NCT00000006", "2027-01-01", ["China"], sponsor_class="OTHER", intervention_types=["DRUG"]),
            study("NCT00000007", "2023-07", ["France"], intervention_types=["PROCEDURE"]),
            study("NCT00000001", "2025-01-01", ["Germany"], sponsor_class="OTHER", intervention_types=["DRUG"]),
        ],
    ]
    files = []
    for index, studies in enumerate(pages, 1):
        path = raw / f"page_{index:06d}.json"
        path.write_text(json.dumps({"studies": studies}), encoding="utf-8")
        files.append({"path": str(path.relative_to(run)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "study_count": len(studies)})
    manifest = {
        "schema_version": 1,
        "harvest_timestamp_utc": "2026-07-16T12:00:00+00:00",
            "request_parameters": {"fields": "NCTId,StartDate,StartDateType,InterventionType,LeadSponsorClass,LocationCountry"},
        "limit_pages": None if complete else 2,
        "page_count": 2,
        "raw_study_count": 8,
        "unique_nct_count": 7,
        "duplicate_nct_count": 1,
        "snapshot_complete": complete,
        "next_page_token_remaining": not complete,
        "files": files,
    }
    (run / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return run


class SummaryOnlyTests(unittest.TestCase):
    def test_standard_and_summary_only_aggregate_parity_and_output_suppression(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            self.assertEqual(main(["--run", str(run), "--output-name", "standard", "--reference-date", "2026-07-16"]), 0)
            self.assertEqual(main(["--run", str(run), "--output-name", "summary", "--reference-date", "2026-07-16", "--summary-only"]), 0)
            standard = json.loads((run / "standard" / "summary.json").read_text())
            summary = json.loads((run / "summary" / "summary.json").read_text())
            for key in ("total_studies", "known_location_studies", "unknown_location_studies", "buckets", "us_presence_groups", "secondary_known_location", "time_analysis"):
                self.assertEqual(standard[key], summary[key], key)
            self.assertEqual(summary["duplicate_nct_ids_detected"], {"NCT00000001": 1})
            self.assertEqual(summary["reconciliation_checks"], standard["reconciliation_checks"])
            self.assertTrue(summary["reconciliation_checks"]["each_time_cohort_geographic_buckets_equal_denominator"])
            self.assertTrue(summary["summary_only"])
            self.assertEqual(summary["analysis_mode"], "SUMMARY_ONLY_STREAMING")
            summary_dir = run / "summary"
            self.assertEqual({path.name for path in summary_dir.iterdir()}, {"summary.json", "summary.md", "nexus_permissible_summary.xlsx"})
            wb = load_workbook(summary_dir / "nexus_permissible_summary.xlsx", read_only=True)
            self.assertEqual(wb.sheetnames, ["Executive_Summary", "Classification_Summary", "Time_Comparison", "Time_Location_Pivot", "Intervention_Type_Audit", "Definitions", "Run_Metadata"])
            self.assertNotIn("Study_Detail", wb.sheetnames)
            self.assertFalse(any(name.startswith("Locations") for name in wb.sheetnames))
            wb.close()
            markdown = (summary_dir / "summary.md").read_text()
            self.assertIn("| RECENT_3Y | NEXUS | 1 |", markdown)
            self.assertIn("| YEARS_4_TO_6_AGO | US_ONLY | 1 |", markdown)
            for cohort, count in summary["time_analysis"]["cohort_counts_all_time"].items():
                self.assertIn(f"| {cohort} | {count:,} |", markdown)
            self.assertIn("| NEXUS % of all cohort |", markdown)
            self.assertIn("| PERMISSIBLE % of known-location |", markdown)

    def test_drug_filter_standard_and_summary_only_aggregate_parity(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            common = ["--run", str(run), "--reference-date", "2026-07-16", "--study-product", "drug"]
            self.assertEqual(main([*common, "--output-name", "standard_drug"]), 0)
            self.assertEqual(main([*common, "--output-name", "summary_drug", "--summary-only"]), 0)
            standard = json.loads((run / "standard_drug" / "summary.json").read_text())
            summary = json.loads((run / "summary_drug" / "summary.json").read_text())
            for key in ("total_studies", "known_location_studies", "buckets", "us_presence_groups",
                        "time_analysis", "intervention_type_audit", "all_vs_drug_diagnostic"):
                self.assertEqual(standard[key], summary[key], key)
            self.assertEqual(summary["study_product_filter"], "drug")
            self.assertEqual(summary["total_studies"], 3)
            self.assertEqual(summary["intervention_type_audit"]["industry_drug_containing_studies"], 3)
            self.assertEqual(summary["intervention_type_audit"]["device_involved_drug_studies"], 1)
            self.assertTrue(summary["reconciliation_checks"]["five_buckets_equal_total"])
            self.assertTrue(summary["reconciliation_checks"]["all_time_cohorts_equal_total"])
            self.assertEqual(summary["time_analysis"]["cohorts"]["RECENT_3Y"]["denominator"], 2)
            self.assertEqual(summary["time_analysis"]["cohorts"]["RECENT_3Y"]["buckets"]["NEXUS"]["count"], 1)
            self.assertEqual(summary["time_analysis"]["cohorts"]["YEARS_4_TO_6_AGO"]["denominator"], 1)
            markdown = (run / "summary_drug" / "summary.md").read_text()
            self.assertIn("Study product filter: **drug**", markdown)
            self.assertIn("## All vs Industry Drug-containing candidate diagnostic comparison", markdown)
            self.assertIn("| DRUG + DEVICE | 1 |", markdown)
            wb = load_workbook(run / "summary_drug" / "nexus_permissible_summary.xlsx", read_only=True)
            self.assertIn("Intervention_Type_Audit", wb.sheetnames)
            wb.close()

    def test_drug_filter_rejects_snapshot_without_required_field(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            manifest_path = run / "manifest.json"
            manifest = json.loads(manifest_path.read_text())
            manifest["request_parameters"]["fields"] = "NCTId,StartDate,LocationCountry"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main(["--run", str(run), "--output-name", "drug", "--summary-only", "--study-product", "drug"])

    def test_summary_only_rejects_hash_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            page = run / "raw" / "page_000001.json"
            page.write_text(page.read_text() + " ", encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "hash verification failed"):
                main(["--run", str(run), "--output-name", "summary", "--summary-only"])

    def test_summary_only_rejects_missing_page(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            (run / "raw" / "page_000002.json").unlink()
            with self.assertRaisesRegex(RuntimeError, "hash verification failed"):
                main(["--run", str(run), "--output-name", "summary", "--summary-only"])

    def test_summary_only_rejects_incomplete_snapshot(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td), complete=False)
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main(["--run", str(run), "--output-name", "summary", "--summary-only"])

    def test_complete_snapshot_rejects_remaining_next_page_token(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            manifest_path = run / "manifest.json"
            manifest = json.loads(manifest_path.read_text())
            manifest["next_page_token_remaining"] = True
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main(["--run", str(run), "--output-name", "summary", "--summary-only"])

    def test_summary_only_rejects_manifest_count_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            manifest_path = run / "manifest.json"
            manifest = json.loads(manifest_path.read_text())
            manifest["raw_study_count"] = 9
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError, "Raw study count disagrees"):
                main(["--run", str(run), "--output-name", "summary", "--summary-only"])

    def test_completed_snapshot_without_start_date_is_explicitly_unavailable(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            manifest_path = run / "manifest.json"
            manifest = json.loads(manifest_path.read_text())
            manifest["request_parameters"]["fields"] = "NCTId,LocationCountry"
            for entry in manifest["files"]:
                page = run / entry["path"]
                payload = json.loads(page.read_text())
                for item in payload["studies"]:
                    item["protocolSection"]["statusModule"].pop("startDateStruct", None)
                page.write_text(json.dumps(payload), encoding="utf-8")
                entry["sha256"] = hashlib.sha256(page.read_bytes()).hexdigest()
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            self.assertEqual(main(["--run", str(run), "--output-name", "summary", "--summary-only"]), 0)
            summary = json.loads((run / "summary" / "summary.json").read_text())
            self.assertEqual(summary["time_analysis_status"], "UNAVAILABLE_START_DATE_NOT_HARVESTED")
            self.assertEqual(summary["time_analysis"]["cohort_counts_all_time"]["TIME_UNKNOWN"], 7)
            self.assertIn("UNAVAILABLE", (run / "summary" / "summary.md").read_text())
            wb = load_workbook(run / "summary" / "nexus_permissible_summary.xlsx", read_only=True)
            pivot = {(row[0], row[1]): row[2:] for row in wb["Time_Location_Pivot"].iter_rows(min_row=2, values_only=True)}
            self.assertIsNone(pivot[("RECENT_3Y", "% of cohort")][-1])
            self.assertIsNone(pivot[("YEARS_4_TO_6_AGO", "% of cohort")][-1])
            wb.close()

    def test_start_date_type_alone_does_not_mark_start_date_available(self):
        with tempfile.TemporaryDirectory() as td:
            run = make_run(Path(td))
            manifest_path = run / "manifest.json"
            manifest = json.loads(manifest_path.read_text())
            manifest["request_parameters"]["fields"] = "NCTId,StartDateType,LocationCountry"
            for entry in manifest["files"]:
                page = run / entry["path"]
                payload = json.loads(page.read_text())
                for item in payload["studies"]:
                    start = item["protocolSection"]["statusModule"].get("startDateStruct")
                    if start:
                        start.pop("date", None)
                page.write_text(json.dumps(payload), encoding="utf-8")
                entry["sha256"] = hashlib.sha256(page.read_bytes()).hexdigest()
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            self.assertEqual(main(["--run", str(run), "--output-name", "summary", "--summary-only"]), 0)
            summary = json.loads((run / "summary" / "summary.json").read_text())
            self.assertFalse(summary["start_date_field_available_in_snapshot"])
            self.assertEqual(summary["time_analysis_status"], "UNAVAILABLE_START_DATE_NOT_HARVESTED")


if __name__ == "__main__":
    unittest.main()
