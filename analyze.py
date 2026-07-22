#!/usr/bin/env python3
"""Analyze an auditable ClinicalTrials.gov harvest snapshot."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import calendar
import math
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BUCKETS = ("NEXUS", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "UNKNOWN")
TRIAL_FIELDS = ["nct_id", "brief_title", "study_type", "overall_status", "start_date_raw", "start_date_parsed",
                "start_date_type", "start_date_precision", "time_cohort", "countries_pipe", "country_count",
                "has_location_data", "has_us", "has_china", "bucket", "us_presence_group", "has_hong_kong", "has_macau",
                "has_taiwan", "has_us_territory"]
LOCATION_FIELDS = ["nct_id", "location_number", "facility", "city", "state", "postal_code", "country"]
TIME_COHORTS = ("RECENT_3Y", "YEARS_4_TO_6_AGO", "OLDER_THAN_6Y", "FUTURE", "TIME_UNKNOWN", "TIME_AMBIGUOUS")
COMPARISON_COHORTS = ("RECENT_3Y", "YEARS_4_TO_6_AGO")
DRUG_CANDIDATE_DEFINITION = "LeadSponsorClass == INDUSTRY AND contains at least one InterventionType == DRUG"
INTERVENTIONAL_DRUG_CANDIDATE_DEFINITION = (
    "StudyType == INTERVENTIONAL AND LeadSponsorClass == INDUSTRY "
    "AND contains at least one InterventionType == DRUG"
)


def sponsor_and_intervention_types(study: dict) -> tuple[str, tuple[str, ...]]:
    protocol = study.get("protocolSection", {})
    sponsor = (protocol.get("sponsorCollaboratorsModule", {}).get("leadSponsor", {}) or {}).get("class")
    interventions = (protocol.get("armsInterventionsModule", {}) or {}).get("interventions") or []
    sponsor_class = sponsor.strip() if isinstance(sponsor, str) else ""
    intervention_types = tuple(sorted({
        value.strip() for item in interventions if isinstance(item, dict)
        for value in [item.get("type")] if isinstance(value, str) and value.strip()
    }, key=lambda value: (value != "DRUG", value)))
    return sponsor_class, intervention_types


def product_is_eligible(study: dict, study_product: str) -> bool:
    if study_product == "all":
        return True
    sponsor_class, intervention_types = sponsor_and_intervention_types(study)
    return sponsor_class == "INDUSTRY" and "DRUG" in intervention_types


def study_type_is_eligible(study: dict, study_type: str) -> bool:
    if study_type == "all":
        return True
    value = (study.get("protocolSection", {}).get("designModule", {}) or {}).get("studyType")
    normalized = value.strip().upper() if isinstance(value, str) else ""
    return normalized == "INTERVENTIONAL"


def update_intervention_type_audit(counts: Counter, combinations: Counter, study: dict) -> None:
    sponsor_class, intervention_types = sponsor_and_intervention_types(study)
    if sponsor_class != "INDUSTRY":
        return
    counts["total_industry_studies"] += 1
    if not intervention_types:
        counts["industry_missing_intervention_type"] += 1
    if "DRUG" not in intervention_types:
        counts["industry_not_drug_containing"] += 1
        return
    counts["industry_drug_containing"] += 1
    combinations[intervention_types] += 1
    if "DEVICE" in intervention_types:
        counts["device_involved_drug_studies"] += 1


def build_intervention_type_audit(counts: Counter, combinations: Counter, study_type: str = "all") -> dict:
    denominator = counts["industry_drug_containing"]
    rows = [
        {"intervention_types": list(values), "combination": " + ".join(values), "unique_study_count": count,
         "percentage_industry_drug_containing": pct(count, denominator)}
        for values, count in sorted(combinations.items(), key=lambda item: (-item[1], item[0]))
    ]
    highlighted = {}
    named = {
        "drug_only": ("DRUG",),
        "drug_other": ("DRUG", "OTHER"),
        "drug_device": ("DRUG", "DEVICE"),
        "drug_procedure": ("DRUG", "PROCEDURE"),
        "drug_behavioral": ("DRUG", "BEHAVIORAL"),
        "drug_dietary_supplement": ("DRUG", "DIETARY_SUPPLEMENT"),
    }
    for key, values in named.items():
        highlighted[key] = combinations[values]
    highlighted["other_observed_combinations_containing_drug"] = denominator - sum(highlighted.values())
    return {
        "lead_sponsor_filter": "INDUSTRY",
        "study_type_filter": study_type,
        "candidate_definition": (INTERVENTIONAL_DRUG_CANDIDATE_DEFINITION
                                 if study_type == "interventional" else DRUG_CANDIDATE_DEFINITION),
        "total_industry_studies": counts["total_industry_studies"],
        "industry_drug_containing_studies": denominator,
        "industry_not_drug_containing_studies": counts["industry_not_drug_containing"],
        "industry_missing_intervention_type": counts["industry_missing_intervention_type"],
        "device_involved_drug_studies": counts["device_involved_drug_studies"],
        "highlighted_combination_counts": highlighted,
        "combination_counts": rows,
    }


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def classify(countries: Iterable[str], cfg: dict) -> tuple[str, dict]:
    """Apply the version-2 hierarchy: UNKNOWN, PERMISSIBLE, US_ONLY, NEXUS, then US_NON_CHINA_MULTI."""
    values = {c.strip() for c in countries if isinstance(c, str) and c.strip()}
    has_us = bool(values & set(cfg["us_countries"]))
    has_china = bool(values & set(cfg["china_countries"]))
    if not values:
        bucket = "UNKNOWN"
        us_presence_group = "UNKNOWN"
    elif not has_us:
        bucket = "PERMISSIBLE"
        us_presence_group = "NO_US"
    elif values == set(cfg["us_countries"]):
        bucket = "US_ONLY"
        us_presence_group = "HAS_US"
    elif has_china:
        bucket = "NEXUS"
        us_presence_group = "HAS_US"
    else:
        bucket = "US_NON_CHINA_MULTI"
        us_presence_group = "HAS_US"
    tracked = cfg["tracked_separately"]
    return bucket, {
        "has_location_data": bool(values), "has_us": has_us, "has_china": has_china,
        "us_presence_group": us_presence_group,
        "has_hong_kong": bool(values & set(tracked["hong_kong"])),
        "has_macau": bool(values & set(tracked["macau"])),
        "has_taiwan": bool(values & set(tracked["taiwan"])),
        "has_us_territory": bool(values & set(tracked["us_territories"])),
    }


def subtract_calendar_years(value: date, years: int) -> date:
    """Subtract whole calendar years, clamping leap day to the target month's final day."""
    target_year = value.year - years
    target_day = min(value.day, calendar.monthrange(target_year, value.month)[1])
    return date(target_year, value.month, target_day)


def parse_start_date(raw: object) -> tuple[str, str, date | None, date | None]:
    """Return canonical value, precision, and conservative possible date interval."""
    if not isinstance(raw, str) or not raw.strip():
        return "", "MISSING", None, None
    value = raw.strip()
    try:
        if len(value) == 10:
            parsed = datetime.strptime(value, "%Y-%m-%d").date()
            return parsed.isoformat(), "DAY", parsed, parsed
        if len(value) == 7:
            parsed = datetime.strptime(value, "%Y-%m").date()
            last = calendar.monthrange(parsed.year, parsed.month)[1]
            return value, "MONTH", date(parsed.year, parsed.month, 1), date(parsed.year, parsed.month, last)
        if len(value) == 4 and value.isdigit():
            year = int(value)
            if year < 1 or year > 9999:
                raise ValueError
            return value, "YEAR", date(year, 1, 1), date(year, 12, 31)
    except ValueError:
        pass
    return "", "UNPARSEABLE", None, None


def assign_time_cohort(raw: object, reference_date: date) -> dict:
    """Assign by interval containment without inventing precision for partial dates."""
    parsed, precision, earliest, latest = parse_start_date(raw)
    if earliest is None or latest is None:
        return {"start_date_parsed": parsed, "start_date_precision": precision, "time_cohort": "TIME_UNKNOWN"}
    three_year_boundary = subtract_calendar_years(reference_date, 3)
    six_year_boundary = subtract_calendar_years(reference_date, 6)
    if earliest > reference_date:
        cohort = "FUTURE"
    elif latest <= six_year_boundary:
        cohort = "OLDER_THAN_6Y"
    elif earliest > three_year_boundary and latest <= reference_date:
        cohort = "RECENT_3Y"
    elif earliest > six_year_boundary and latest <= three_year_boundary:
        cohort = "YEARS_4_TO_6_AGO"
    else:
        cohort = "TIME_AMBIGUOUS"
    return {"start_date_parsed": parsed, "start_date_precision": precision, "time_cohort": cohort}


def build_time_analysis_from_counts(time_bucket_counts: dict[str, Counter], reference_date: date, expected_total: int | None = None) -> dict:
    cohort_counts = {}
    cohort_reconciliation = {}
    for cohort in TIME_COHORTS:
        counts = time_bucket_counts[cohort]
        unexpected = {bucket: count for bucket, count in counts.items() if bucket not in BUCKETS and count}
        if unexpected:
            raise AssertionError(f"{cohort} contains unexpected geographic buckets: {unexpected}")
        denominator = sum(counts.values())
        five_bucket_total = sum(counts[bucket] for bucket in BUCKETS)
        cohort_counts[cohort] = denominator
        cohort_reconciliation[cohort] = five_bucket_total - denominator
        if cohort_reconciliation[cohort] != 0:
            raise AssertionError(f"{cohort} geographic buckets do not reconcile")
    cohort_total = sum(cohort_counts.values())
    if cohort_total <= 0:
        raise AssertionError("Time cohorts contain no studies")
    if expected_total is not None and cohort_total != expected_total:
        raise AssertionError("All time cohorts do not reconcile to expected total")
    cohorts = {}
    for cohort in COMPARISON_COHORTS:
        counts = time_bucket_counts[cohort]
        total = cohort_counts[cohort]
        known = total - counts["UNKNOWN"]
        has_us = counts["US_ONLY"] + counts["US_NON_CHINA_MULTI"] + counts["NEXUS"]
        if sum(counts[b] for b in BUCKETS) != total:
            raise AssertionError(f"{cohort} geographic buckets do not reconcile")
        cohorts[cohort] = {
            "denominator": total,
            "known_location_denominator": known,
            "buckets": {
                b: {"count": counts[b], "percentage_cohort": pct(counts[b], total),
                    "percentage_known_locations": None if b == "UNKNOWN" else pct(counts[b], known)}
                for b in BUCKETS
            },
            "has_us": {"count": has_us, "percentage_cohort": pct(has_us, total),
                       "percentage_known_locations": pct(has_us, known)},
            "reconciliation_difference": sum(counts[b] for b in BUCKETS) - total,
        }
    metrics = [
        ("NEXUS count", "count", "NEXUS"),
        ("NEXUS % of all cohort", "percentage_cohort", "NEXUS"),
        ("NEXUS % of known-location", "percentage_known_locations", "NEXUS"),
        ("PERMISSIBLE count", "count", "PERMISSIBLE"),
        ("PERMISSIBLE % of all cohort", "percentage_cohort", "PERMISSIBLE"),
        ("PERMISSIBLE % of known-location", "percentage_known_locations", "PERMISSIBLE"),
        ("US_ONLY % of all cohort", "percentage_cohort", "US_ONLY"),
        ("US_NON_CHINA_MULTI % of all cohort", "percentage_cohort", "US_NON_CHINA_MULTI"),
        ("UNKNOWN % of all cohort", "percentage_cohort", "UNKNOWN"),
    ]
    comparison = [
        {"metric": "Cohort denominator", "recent_3y": cohorts["RECENT_3Y"]["denominator"],
         "years_4_to_6_ago": cohorts["YEARS_4_TO_6_AGO"]["denominator"],
         "change": cohorts["RECENT_3Y"]["denominator"] - cohorts["YEARS_4_TO_6_AGO"]["denominator"], "change_unit": "count"},
        {"metric": "Known-location denominator", "recent_3y": cohorts["RECENT_3Y"]["known_location_denominator"],
         "years_4_to_6_ago": cohorts["YEARS_4_TO_6_AGO"]["known_location_denominator"],
         "change": cohorts["RECENT_3Y"]["known_location_denominator"] - cohorts["YEARS_4_TO_6_AGO"]["known_location_denominator"], "change_unit": "count"},
    ]
    for label, field, bucket in metrics:
        recent = cohorts["RECENT_3Y"]["buckets"][bucket][field]
        older = cohorts["YEARS_4_TO_6_AGO"]["buckets"][bucket][field]
        comparison.append({"metric": label, "recent_3y": recent, "years_4_to_6_ago": older,
                           "change": recent - older if recent is not None and older is not None else None,
                           "change_unit": "count" if field == "count" else "percentage_points"})
    for label, field in (("HAS_US count", "count"), ("HAS_US % of all cohort", "percentage_cohort"),
                         ("HAS_US % of known-location", "percentage_known_locations")):
        recent = cohorts["RECENT_3Y"]["has_us"][field]
        older = cohorts["YEARS_4_TO_6_AGO"]["has_us"][field]
        comparison.append({"metric": label, "recent_3y": recent, "years_4_to_6_ago": older,
                           "change": recent - older if recent is not None and older is not None else None,
                           "change_unit": "count" if field == "count" else "percentage_points"})
    three = subtract_calendar_years(reference_date, 3)
    six = subtract_calendar_years(reference_date, 6)
    return {
        "temporal_field": "registered Start Date",
        "reference_date": reference_date.isoformat(),
        "boundaries": {
            "RECENT_3Y": {"start_exclusive": three.isoformat(), "end_inclusive": reference_date.isoformat()},
            "YEARS_4_TO_6_AGO": {"start_exclusive": six.isoformat(), "end_inclusive": three.isoformat()},
        },
        "partial_date_rule": "Assign only when the entire possible date interval is contained in one cohort; otherwise TIME_AMBIGUOUS.",
        "cohort_counts_all_time": cohort_counts,
        "time_unknown_count": cohort_counts["TIME_UNKNOWN"],
        "time_ambiguous_count": cohort_counts["TIME_AMBIGUOUS"],
        "all_time_cohort_reconciliation_difference": 0 if expected_total is None else cohort_total - expected_total,
        "cohort_geographic_reconciliation_differences": cohort_reconciliation,
        "time_location_counts": {cohort: {bucket: time_bucket_counts[cohort][bucket] for bucket in BUCKETS} for cohort in TIME_COHORTS},
        "cohorts": cohorts,
        "comparison": comparison,
    }


def build_time_analysis(trials: list[dict], reference_date: date) -> dict:
    time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
    for row in trials:
        time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
    result = build_time_analysis_from_counts(time_bucket_counts, reference_date, len(trials))
    if sum(result["cohort_counts_all_time"].values()) != len(trials):
        raise AssertionError("All time cohorts do not reconcile to total studies")
    return result


def build_all_vs_drug_diagnostic(all_time: dict, drug_time: dict) -> list[dict]:
    rows = []
    for cohort, label in (("RECENT_3Y", "Recent 3Y"), ("YEARS_4_TO_6_AGO", "4–6Y")):
        all_item = all_time["cohorts"][cohort]
        drug_item = drug_time["cohorts"][cohort]
        rows.extend([
            {"metric": f"{label} denominator", "all_studies": all_item["denominator"],
             "drug_studies": drug_item["denominator"], "value_type": "count"},
            {"metric": f"{label} Nexus %", "all_studies": all_item["buckets"]["NEXUS"]["percentage_cohort"],
             "drug_studies": drug_item["buckets"]["NEXUS"]["percentage_cohort"], "value_type": "percentage"},
            {"metric": f"{label} Permissible %", "all_studies": all_item["buckets"]["PERMISSIBLE"]["percentage_cohort"],
             "drug_studies": drug_item["buckets"]["PERMISSIBLE"]["percentage_cohort"], "value_type": "percentage"},
        ])
    return rows


def extract(study: dict, cfg: dict, reference_date: date | None = None, include_locations: bool = True) -> tuple[dict, list[dict]]:
    protocol = study.get("protocolSection", {})
    ident = protocol.get("identificationModule", {})
    design = protocol.get("designModule", {})
    status = protocol.get("statusModule", {})
    contacts = protocol.get("contactsLocationsModule", {})
    nct = ident.get("nctId")
    if not isinstance(nct, str) or not nct.strip():
        raise ValueError("Every retained record must have a non-empty NCT ID")
    locations = []
    countries = set()
    for number, loc in enumerate(contacts.get("locations") or [], 1):
        country = loc.get("country")
        if isinstance(country, str) and country.strip():
            countries.add(country.strip())
        if include_locations:
            locations.append({"nct_id": nct, "location_number": number, "facility": loc.get("facility", ""),
                              "city": loc.get("city", ""), "state": loc.get("state", ""),
                              "postal_code": loc.get("zip", ""), "country": country or ""})
    bucket, flags = classify(countries, cfg)
    start_struct = status.get("startDateStruct") or {}
    start_raw = start_struct.get("date", "")
    time_fields = assign_time_cohort(start_raw, reference_date) if reference_date else {
        "start_date_parsed": parse_start_date(start_raw)[0],
        "start_date_precision": parse_start_date(start_raw)[1],
        "time_cohort": "TIME_UNKNOWN",
    }
    row = {"nct_id": nct, "brief_title": ident.get("briefTitle", ""), "study_type": design.get("studyType", ""),
           "overall_status": status.get("overallStatus", ""), "start_date_raw": start_raw,
           "start_date_type": start_struct.get("type", ""), **time_fields,
           "countries_pipe": "|".join(sorted(countries)),
           "country_count": len(countries), **flags, "bucket": bucket}
    return row, locations


def pct(n: int, d: int) -> float | None:
    return n / d if d else None


def validate_trial_row(row: dict, cfg: dict) -> None:
    bucket = row["bucket"]
    countries = set(filter(None, row["countries_pipe"].split("|")))
    if bucket == "NEXUS" and not (row["has_us"] and row["has_china"]):
        raise AssertionError("Invalid NEXUS")
    if bucket == "US_ONLY" and countries != set(cfg["us_countries"]):
        raise AssertionError("Invalid US_ONLY")
    if bucket == "US_NON_CHINA_MULTI" and not (row["has_us"] and not row["has_china"] and len(countries) > 1):
        raise AssertionError("Invalid US_NON_CHINA_MULTI")
    if bucket == "PERMISSIBLE" and not (row["has_location_data"] and not row["has_us"]):
        raise AssertionError("Invalid PERMISSIBLE")
    if bucket == "UNKNOWN" and row["has_location_data"]:
        raise AssertionError("Invalid UNKNOWN")
    expected_group = "UNKNOWN" if bucket == "UNKNOWN" else "NO_US" if bucket == "PERMISSIBLE" else "HAS_US"
    if row["us_presence_group"] != expected_group:
        raise AssertionError("Invalid us_presence_group")
    if row["time_cohort"] not in TIME_COHORTS:
        raise AssertionError("Invalid time_cohort")


def build_summary(*, counts: Counter, total: int, time_analysis: dict, manifest: dict, cfg: dict,
                  duplicate_ids: Counter, reference_date: date, run_directory: Path,
                  summary_only: bool, analysis_timestamp_utc: str, study_product: str, study_type: str,
                  intervention_type_audit: dict | None, all_vs_drug_diagnostic: list[dict] | None) -> dict:
    known = total - counts["UNKNOWN"]
    has_us_count = counts["US_ONLY"] + counts["US_NON_CHINA_MULTI"] + counts["NEXUS"]
    if sum(counts[b] for b in BUCKETS) != total:
        raise AssertionError("Buckets do not sum to denominator")
    if counts["PERMISSIBLE"] + has_us_count != known:
        raise AssertionError("NO_US + HAS_US does not equal known-location studies")
    if sum(time_analysis["cohort_counts_all_time"].values()) != total:
        raise AssertionError("All time cohorts do not reconcile to total unique studies")
    if any(time_analysis["cohort_geographic_reconciliation_differences"].values()):
        raise AssertionError("One or more time cohorts do not reconcile across the five geographic buckets")
    for cohort in COMPARISON_COHORTS:
        if time_analysis["cohorts"][cohort]["reconciliation_difference"] != 0:
            raise AssertionError(f"{cohort} reconciliation difference is nonzero")
    bucket_items = {}
    for bucket in BUCKETS:
        bucket_items[bucket] = {
            "count": counts[bucket],
            "percentage_total": pct(counts[bucket], total),
            "percentage_known_locations": pct(counts[bucket], known) if bucket != "UNKNOWN" else None,
            "percentage_has_us": pct(counts[bucket], has_us_count) if bucket in {"US_ONLY", "US_NON_CHINA_MULTI", "NEXUS"} else None,
        }
    requested_fields = {
        field.strip() for field in manifest.get("request_parameters", {}).get("fields", "").split(",") if field.strip()
    }
    start_date_available = "StartDate" in requested_fields
    intervention_type_available = "InterventionType" in requested_fields
    lead_sponsor_class_available = "LeadSponsorClass" in requested_fields
    return {
        "result_status": "FINAL_COMPLETE_SNAPSHOT" if manifest["snapshot_complete"] else "PARTIAL_NON_FINAL_SMOKE_TEST",
        "snapshot_complete": manifest["snapshot_complete"],
        "harvest_timestamp_utc": manifest["harvest_timestamp_utc"],
        "analysis_timestamp_utc": analysis_timestamp_utc,
        "analysis_mode": "SUMMARY_ONLY_STREAMING" if summary_only else "STANDARD_DETAIL",
        "summary_only": summary_only,
        "run_directory": str(run_directory.resolve()),
        "raw_hash_verification": "PASS",
        "classification_model_version": 2,
        "reference_date": reference_date.isoformat(),
        "start_date_field_available_in_snapshot": start_date_available,
        "time_analysis_status": "AVAILABLE" if start_date_available else "UNAVAILABLE_START_DATE_NOT_HARVESTED",
        "unit_of_analysis": "one unique NCT ID",
        "study_product_filter": study_product,
        "study_type_filter": study_type,
        "drug_candidate_definition": (INTERVENTIONAL_DRUG_CANDIDATE_DEFINITION
                                      if study_type == "interventional" and study_product == "drug"
                                      else DRUG_CANDIDATE_DEFINITION),
        "final_intervention_exclusion_rule": "NOT_YET_FIXED_PENDING_COMBINATION_AUDIT",
        "intervention_type_field_available_in_snapshot": intervention_type_available,
        "lead_sponsor_class_field_available_in_snapshot": lead_sponsor_class_available,
        "snapshot_unique_studies": manifest["unique_nct_count"],
        "intervention_type_audit": intervention_type_audit,
        "all_vs_drug_diagnostic": all_vs_drug_diagnostic,
        "total_studies": total,
        "known_location_studies": known,
        "unknown_location_studies": counts["UNKNOWN"],
        "buckets": bucket_items,
        "us_presence_groups": {
            "UNKNOWN": {"count": counts["UNKNOWN"], "percentage_total": pct(counts["UNKNOWN"], total), "percentage_known_locations": None},
            "NO_US": {"count": counts["PERMISSIBLE"], "percentage_total": pct(counts["PERMISSIBLE"], total), "percentage_known_locations": pct(counts["PERMISSIBLE"], known)},
            "HAS_US": {"count": has_us_count, "percentage_total": pct(has_us_count, total), "percentage_known_locations": pct(has_us_count, known)},
        },
        "reconciliation_checks": {
            "five_buckets_equal_total": sum(counts[b] for b in BUCKETS) == total,
            "no_us_plus_has_us_equal_known": counts["PERMISSIBLE"] + has_us_count == known,
            "us_subcategories_equal_has_us": counts["US_ONLY"] + counts["US_NON_CHINA_MULTI"] + counts["NEXUS"] == has_us_count,
            "all_time_cohorts_equal_total": sum(time_analysis["cohort_counts_all_time"].values()) == total,
            "each_time_cohort_geographic_buckets_equal_denominator": not any(
                time_analysis["cohort_geographic_reconciliation_differences"].values()
            ),
            "primary_time_cohorts_reconcile": all(time_analysis["cohorts"][cohort]["reconciliation_difference"] == 0 for cohort in COMPARISON_COHORTS),
        },
        "secondary_known_location": {"denominator": known, "nexus_percentage": pct(counts["NEXUS"], known),
                                     "permissible_percentage": pct(counts["PERMISSIBLE"], known)},
        "definitions": {"us_countries": cfg["us_countries"], "china_countries": cfg["china_countries"],
                        "tracked_separately": cfg["tracked_separately"]},
        "duplicate_nct_ids_detected": dict(duplicate_ids),
        "time_analysis": time_analysis,
        "location_proxy_limitation": "ClinicalTrials.gov locations are registered or planned facilities; they do not confirm participant nationality or actual enrollment by country.",
    }


def write_csv(path: Path, fields: list[str], rows: Iterable[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_summary_files(out: Path, summary: dict, cfg: dict, manifest: dict) -> None:
    (out / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    total = summary["total_studies"]
    known = summary["known_location_studies"]
    counts = {bucket: summary["buckets"][bucket]["count"] for bucket in BUCKETS}
    has_us_count = summary["us_presence_groups"]["HAS_US"]["count"]
    time_analysis = summary["time_analysis"]
    reference_date = summary["reference_date"]
    label = "FINAL — COMPLETE SNAPSHOT" if manifest["snapshot_complete"] else "PARTIAL / NON-FINAL SMOKE TEST"
    fmt = lambda value: "N/A" if value is None else f"{value:.2%}"
    lines = [f"# Nexus and Permissible Analysis — Five-Category Model — {label}", "", "## Overall population", "",
             f"- Analysis mode: **{summary['analysis_mode']}**", f"- Study type filter: **{summary['study_type_filter']}**",
             f"- Study product filter: **{summary['study_product_filter']}**",
             f"- Drug candidate definition: **{summary['drug_candidate_definition']}**",
             f"- Total studies: **{total:,}**", f"- Known-location studies: **{known:,}**", f"- UNKNOWN: **{counts['UNKNOWN']:,} ({fmt(pct(counts['UNKNOWN'], total))} of all studies)**", "",
             "## Core four-category summary", "",
             "NON_EXEMPT combines US_ONLY and US_NON_CHINA_MULTI. Percentages use the complete analysis denominator.", "",
             "| Category | Study count | Percentage |", "|---|---:|---:|",
             f"| NEXUS | {counts['NEXUS']:,} | {fmt(pct(counts['NEXUS'], total))} |",
             f"| PERMISSIBLE | {counts['PERMISSIBLE']:,} | {fmt(pct(counts['PERMISSIBLE'], total))} |",
             f"| NON_EXEMPT | {counts['US_ONLY'] + counts['US_NON_CHINA_MULTI']:,} | {fmt(pct(counts['US_ONLY'] + counts['US_NON_CHINA_MULTI'], total))} |",
             f"| UNKNOWN | {counts['UNKNOWN']:,} | {fmt(pct(counts['UNKNOWN'], total))} |", "",
             "## High-level geographic split", "", f"- NO_US / PERMISSIBLE: **{counts['PERMISSIBLE']:,} ({fmt(pct(counts['PERMISSIBLE'], total))} of all; {fmt(pct(counts['PERMISSIBLE'], known))} of known-location)**",
             f"- HAS_US: **{has_us_count:,} ({fmt(pct(has_us_count, total))} of all; {fmt(pct(has_us_count, known))} of known-location)**", "",
             "## Breakdown of HAS_US studies", "",
             f"- US_ONLY: **{counts['US_ONLY']:,} ({fmt(pct(counts['US_ONLY'], total))} of all; {fmt(pct(counts['US_ONLY'], known))} of known-location; {fmt(pct(counts['US_ONLY'], has_us_count))} of HAS_US)**",
             f"- US_NON_CHINA_MULTI: **{counts['US_NON_CHINA_MULTI']:,} ({fmt(pct(counts['US_NON_CHINA_MULTI'], total))} of all; {fmt(pct(counts['US_NON_CHINA_MULTI'], known))} of known-location; {fmt(pct(counts['US_NON_CHINA_MULTI'], has_us_count))} of HAS_US)**",
             f"- NEXUS: **{counts['NEXUS']:,} ({fmt(pct(counts['NEXUS'], total))} of all; {fmt(pct(counts['NEXUS'], known))} of known-location; {fmt(pct(counts['NEXUS'], has_us_count))} of HAS_US)**", "",
             "## Reconciliation checks", "", f"- Five buckets equal total: **{sum(counts[b] for b in BUCKETS):,} = {total:,}**",
             f"- PERMISSIBLE + HAS_US = known-location: **{counts['PERMISSIBLE']:,} + {has_us_count:,} = {known:,}**",
             f"- US_ONLY + US_NON_CHINA_MULTI + NEXUS = HAS_US: **{counts['US_ONLY']:,} + {counts['US_NON_CHINA_MULTI']:,} + {counts['NEXUS']:,} = {has_us_count:,}**", ""]
    if summary["time_analysis_status"] != "AVAILABLE":
        lines += ["## Start-date period comparison", "",
                  "**UNAVAILABLE:** this snapshot did not harvest StartDate. All-time geographic results remain valid, but time-period results must not be reported.", ""]
    else:
        lines += ["## Start-date period comparison", "", f"- Reference date: **{reference_date}**",
                  f"- RECENT_3Y: **{time_analysis['boundaries']['RECENT_3Y']['start_exclusive']} < Start Date <= {time_analysis['boundaries']['RECENT_3Y']['end_inclusive']}**",
                  f"- YEARS_4_TO_6_AGO: **{time_analysis['boundaries']['YEARS_4_TO_6_AGO']['start_exclusive']} < Start Date <= {time_analysis['boundaries']['YEARS_4_TO_6_AGO']['end_inclusive']}**",
                  f"- TIME_UNKNOWN: **{time_analysis['time_unknown_count']:,}**; TIME_AMBIGUOUS: **{time_analysis['time_ambiguous_count']:,}**", "",
                  "### All time-cohort counts", "", "| Time cohort | Count |", "|---|---:|"]
        for cohort in TIME_COHORTS:
            lines.append(f"| {cohort} | {time_analysis['cohort_counts_all_time'][cohort]:,} |")
        lines += ["",
                  "| Cohort | Denominator | Known location | NEXUS | NEXUS % cohort | NEXUS % known | PERMISSIBLE | PERMISSIBLE % cohort | PERMISSIBLE % known |",
                  "|---|---:|---:|---:|---:|---:|---:|---:|---:|"]
        for cohort in COMPARISON_COHORTS:
            item = time_analysis["cohorts"][cohort]
            nexus = item["buckets"]["NEXUS"]
            permissible = item["buckets"]["PERMISSIBLE"]
            lines.append(f"| {cohort} | {item['denominator']:,} | {item['known_location_denominator']:,} | {nexus['count']:,} | {fmt(nexus['percentage_cohort'])} | {fmt(nexus['percentage_known_locations'])} | {permissible['count']:,} | {fmt(permissible['percentage_cohort'])} | {fmt(permissible['percentage_known_locations'])} |")
        lines += ["", "### Full geographic distribution within primary cohorts", "",
                  "| Cohort | Bucket | Count | % of cohort | % of known-location cohort |",
                  "|---|---|---:|---:|---:|"]
        for cohort in COMPARISON_COHORTS:
            item = time_analysis["cohorts"][cohort]
            for bucket in BUCKETS:
                bucket_item = item["buckets"][bucket]
                lines.append(
                    f"| {cohort} | {bucket} | {bucket_item['count']:,} | "
                    f"{fmt(bucket_item['percentage_cohort'])} | {fmt(bucket_item['percentage_known_locations'])} |"
                )
            lines.append(
                f"| {cohort} | HAS_US | {item['has_us']['count']:,} | "
                f"{fmt(item['has_us']['percentage_cohort'])} | {fmt(item['has_us']['percentage_known_locations'])} |"
            )
        lines += ["", "### Direct Recent 3Y vs 4–6 Years Ago comparison", "",
                  "| Metric | Recent 3Y | 4–6 Years Ago | Change | Change unit |",
                  "|---|---:|---:|---:|---|"]
        for comparison_row in time_analysis["comparison"]:
            is_percentage = comparison_row["change_unit"] == "percentage_points"
            recent_value = fmt(comparison_row["recent_3y"]) if is_percentage else f"{comparison_row['recent_3y']:,}"
            older_value = fmt(comparison_row["years_4_to_6_ago"]) if is_percentage else f"{comparison_row['years_4_to_6_ago']:,}"
            change_value = fmt(comparison_row["change"]) if is_percentage else f"{comparison_row['change']:+,}"
            lines.append(
                f"| {comparison_row['metric']} | {recent_value} | {older_value} | {change_value} | "
                f"{'percentage points' if is_percentage else 'count'} |"
            )
        lines += ["", "Registered Start Date is a registry field, not confirmed first-participant enrollment. Period comparisons are descriptive and do not imply causation.", ""]
    audit = summary["intervention_type_audit"]
    if audit is not None:
        lines += ["## Industry Drug-containing intervention-type audit", "",
                  f"- Lead Sponsor filter: **{audit['lead_sponsor_filter']}**",
                  f"- Total Industry studies: **{audit['total_industry_studies']:,}**",
                  f"- Industry Drug-containing studies: **{audit['industry_drug_containing_studies']:,}**",
                  f"- Industry studies not containing DRUG: **{audit['industry_not_drug_containing_studies']:,}**",
                  f"- Industry studies with missing intervention type: **{audit['industry_missing_intervention_type']:,}**",
                  f"- Device-involved Industry Drug studies: **{audit['device_involved_drug_studies']:,}**", "",
                  "| Intervention Type Combination | Unique Study Count | % of Industry Drug-containing Studies |",
                  "|---|---:|---:|"]
        for row in audit["combination_counts"]:
            lines.append(f"| {row['combination']} | {row['unique_study_count']:,} | {fmt(row['percentage_industry_drug_containing'])} |")
        highlight_labels = {
            "drug_only": "DRUG only", "drug_other": "DRUG + OTHER", "drug_device": "DRUG + DEVICE",
            "drug_procedure": "DRUG + PROCEDURE", "drug_behavioral": "DRUG + BEHAVIORAL",
            "drug_dietary_supplement": "DRUG + DIETARY_SUPPLEMENT",
            "other_observed_combinations_containing_drug": "Other observed combinations containing DRUG",
        }
        lines += ["", "### Highlighted combination counts", "", "| Category | Unique Study Count |", "|---|---:|"]
        for key, label_text in highlight_labels.items():
            lines.append(f"| {label_text} | {audit['highlighted_combination_counts'][key]:,} |")
        lines += ["", "The final exclusion rule for non-drug intervention combinations has NOT yet been fixed. It will be decided after reviewing the actual intervention-type combination distribution.", ""]
    if summary["all_vs_drug_diagnostic"] is not None:
        comparison_population = ("All INTERVENTIONAL Studies" if summary["study_type_filter"] == "interventional"
                                 else "All Studies")
        lines += ["## All vs Industry Drug-containing candidate diagnostic comparison", "", f"| Metric | {comparison_population} | Industry Drug-containing Candidates |",
                  "|---|---:|---:|"]
        for row in summary["all_vs_drug_diagnostic"]:
            all_value = fmt(row["all_studies"]) if row["value_type"] == "percentage" else f"{row['all_studies']:,}"
            drug_value = fmt(row["drug_studies"]) if row["value_type"] == "percentage" else f"{row['drug_studies']:,}"
            lines.append(f"| {row['metric']} | {all_value} | {drug_value} |")
        lines.append("")
    lines += ["## Definitions and limitations", "",
              "- US_ONLY means the complete unique country set is exactly {United States}.",
              "- US_NON_CHINA_MULTI contains United States plus at least one other non-China country.",
              "- NEXUS contains both United States and China and may contain other countries.",
              "- PERMISSIBLE has a known country set with no United States; NEXUS and PERMISSIBLE alone do not partition the registry.",
              f"- US definition: {', '.join(cfg['us_countries'])}", f"- China definition: {', '.join(cfg['china_countries'])}",
              "- Unit of analysis: one unique NCT ID", f"- Harvest timestamp: {manifest['harvest_timestamp_utc']}", f"- Analysis timestamp: {summary['analysis_timestamp_utc']}",
              f"- Study type filter: {summary['study_type_filter']}",
              f"- Denominator: unique NCT IDs after applying study_type_filter={summary['study_type_filter']} and study_product_filter={summary['study_product_filter']}." if manifest["snapshot_complete"] else "- This page-limited snapshot is not the target denominator and must not be reported as final.",
              "- Location-proxy limitation: registered or planned facilities do not confirm participant nationality or actual country-level enrollment."]
    (out / "summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def prepare_stream_sheet(ws, headers: list[str], widths: list[int] | None = None):
    """Prepare a write-only worksheet with a styled, frozen, filterable header."""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    cells = []
    for value in headers:
        cell = WriteOnlyCell(ws, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)
        cells.append(cell)
    ws.append(cells)
    for index, width in enumerate(widths or [18] * len(headers), 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def append_values(ws, values, percentage_columns=()):
    cells = []
    for index, value in enumerate(values, 1):
        cell = WriteOnlyCell(ws, value=value)
        if index in percentage_columns:
            cell.number_format = "0.00%"
        cells.append(cell)
    ws.append(cells)


def add_intervention_type_audit_sheet(wb: Workbook, audit: dict | None) -> None:
    if audit is None:
        return
    ws = wb.create_sheet("Intervention_Type_Audit")
    prepare_stream_sheet(ws, ["Intervention Type Combination", "Unique Study Count", "% of Industry Drug-containing Studies"], [48, 22, 38])
    append_values(ws, ["Lead Sponsor filter", audit["lead_sponsor_filter"], None])
    append_values(ws, ["Total Industry studies", audit["total_industry_studies"], None])
    append_values(ws, ["Industry Drug-containing studies", audit["industry_drug_containing_studies"], None])
    append_values(ws, ["Industry studies not containing DRUG", audit["industry_not_drug_containing_studies"], None])
    append_values(ws, ["Industry studies with missing intervention type", audit["industry_missing_intervention_type"], None])
    append_values(ws, ["Device-involved Industry Drug studies", audit["device_involved_drug_studies"], None])
    append_values(ws, ["EXACT OBSERVED COMBINATIONS", None, None])
    for row in audit["combination_counts"]:
        append_values(ws, [row["combination"], row["unique_study_count"], row["percentage_industry_drug_containing"]], (3,))
    append_values(ws, ["HIGHLIGHTED EXACT COMBINATIONS", None, None])
    for key, count in audit["highlighted_combination_counts"].items():
        append_values(ws, [key, count, pct(count, audit["industry_drug_containing_studies"])], (3,))


def make_workbook(path: Path, summary: dict, trials: list[dict], locations: list[dict], country_rows: list[dict], cfg: dict, manifest: dict, location_rows_per_sheet: int = 500_000):
    if not 1 <= location_rows_per_sheet <= 1_048_575:
        raise ValueError("location_rows_per_sheet must leave room for the Excel header row")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Executive_Summary")
    prepare_stream_sheet(ws, ["Nexus & Permissible Study Location Analysis — Five-Category Model", "Value"], [58, 95])
    append_values(ws, ["Result status", summary["result_status"]])
    append_values(ws, ["Study type filter", summary["study_type_filter"]])
    append_values(ws, ["Study product filter", summary["study_product_filter"]])
    append_values(ws, ["Drug candidate definition", summary["drug_candidate_definition"]])
    append_values(ws, ["Denominator", "Unique Industry-sponsored DRUG-containing NCT IDs." if summary["study_product_filter"] == "drug" else "All unique study records accessible through the ClinicalTrials.gov API in this completed snapshot."])
    append_values(ws, ["Snapshot applicability", "Completed all-study snapshot; primary denominator applies." if summary["snapshot_complete"] else "PARTIAL smoke-test snapshot; the target denominator above is not satisfied and these results are not final."])
    append_values(ws, ["Location interpretation", "Registered or planned study facilities, not confirmed participant nationality or actual country-level enrollment."])
    append_values(ws, ["OVERALL POPULATION", None])
    append_values(ws, ["Total unique studies", summary["total_studies"]])
    append_values(ws, ["Known-location studies", summary["known_location_studies"]])
    append_values(ws, ["Unknown-location studies", summary["buckets"]["UNKNOWN"]["count"]])
    append_values(ws, ["HIGH-LEVEL GEOGRAPHIC SPLIT", None])
    append_values(ws, ["NO_US / PERMISSIBLE count", summary["us_presence_groups"]["NO_US"]["count"]])
    append_values(ws, ["NO_US / PERMISSIBLE % of all studies", summary["us_presence_groups"]["NO_US"]["percentage_total"]], (2,))
    append_values(ws, ["NO_US / PERMISSIBLE % of known-location studies", summary["us_presence_groups"]["NO_US"]["percentage_known_locations"]], (2,))
    append_values(ws, ["HAS_US count", summary["us_presence_groups"]["HAS_US"]["count"]])
    append_values(ws, ["HAS_US % of all studies", summary["us_presence_groups"]["HAS_US"]["percentage_total"]], (2,))
    append_values(ws, ["HAS_US % of known-location studies", summary["us_presence_groups"]["HAS_US"]["percentage_known_locations"]], (2,))
    append_values(ws, ["BREAKDOWN OF HAS_US STUDIES", None])
    for b in ("US_ONLY", "US_NON_CHINA_MULTI", "NEXUS"):
        append_values(ws, [f"{b} count", summary["buckets"][b]["count"]])
        append_values(ws, [f"{b} % of all studies", summary["buckets"][b]["percentage_total"]], (2,))
        append_values(ws, [f"{b} % of known-location studies", summary["buckets"][b]["percentage_known_locations"]], (2,))
        append_values(ws, [f"{b} % of HAS_US studies", summary["buckets"][b]["percentage_has_us"]], (2,))
    append_values(ws, ["START-DATE PERIOD COMPARISON", None])
    append_values(ws, ["Reference date", summary["time_analysis"]["reference_date"]])
    for cohort, label in (("RECENT_3Y", "Recent 3 years"), ("YEARS_4_TO_6_AGO", "4–6 years ago")):
        item = summary["time_analysis"]["cohorts"][cohort]
        append_values(ws, [f"{label} denominator", item["denominator"]])
        for bucket in ("NEXUS", "PERMISSIBLE"):
            append_values(ws, [f"{label} {bucket} count", item["buckets"][bucket]["count"]])
            append_values(ws, [f"{label} {bucket} % of cohort", item["buckets"][bucket]["percentage_cohort"]], (2,))
    if summary["all_vs_drug_diagnostic"] is not None:
        append_values(ws, ["ALL VS DRUG DIAGNOSTIC", None])
        for row in summary["all_vs_drug_diagnostic"]:
            append_values(ws, [f"{row['metric']} — All Studies", row["all_studies"]], (2,) if row["value_type"] == "percentage" else ())
            append_values(ws, [f"{row['metric']} — Industry Drug-containing Candidates", row["drug_studies"]], (2,) if row["value_type"] == "percentage" else ())

    cs = wb.create_sheet("Classification_Summary")
    prepare_stream_sheet(cs, ["bucket", "count", "percentage_total", "percentage_known_locations", "percentage_has_us", "denominator_notes"], [24, 14, 20, 27, 22, 55])
    for b in BUCKETS:
        item = summary["buckets"][b]
        append_values(cs, [b, item["count"], item["percentage_total"], item["percentage_known_locations"], item["percentage_has_us"],
                           "Primary: all studies; known-location and HAS_US percentages are secondary diagnostics"], (3, 4, 5))

    tc = wb.create_sheet("Time_Comparison")
    prepare_stream_sheet(tc, ["Metric", "Recent 3 years", "4–6 years ago", "Change", "Change unit / denominator"], [38, 22, 22, 22, 52])
    for row in summary["time_analysis"]["comparison"]:
        is_percentage = row["change_unit"] == "percentage_points"
        append_values(tc, [row["metric"], row["recent_3y"], row["years_4_to_6_ago"], row["change"],
                           "Recent minus 4–6 years; percentage-point change" if is_percentage else "Recent count minus 4–6 years count"],
                      (2, 3, 4) if is_percentage else ())

    pivot = wb.create_sheet("Time_Location_Pivot")
    pivot_buckets = ["NEXUS", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "UNKNOWN"]
    prepare_stream_sheet(pivot, ["Time cohort", "Measure", *pivot_buckets, "Total"], [25, 18, 16, 18, 16, 28, 16, 16])
    for cohort in TIME_COHORTS:
        counts = summary["time_analysis"]["time_location_counts"][cohort]
        denominator = summary["time_analysis"]["cohort_counts_all_time"][cohort]
        append_values(pivot, [cohort, "Count", *[counts[b] for b in pivot_buckets], denominator])
        if cohort in COMPARISON_COHORTS:
            item = summary["time_analysis"]["cohorts"][cohort]
            total_percentage = 1.0 if denominator else None
            append_values(pivot, [cohort, "% of cohort", *[item["buckets"][b]["percentage_cohort"] for b in pivot_buckets], total_percentage], tuple(range(3, 9)))

    detail = wb.create_sheet("Study_Detail")
    prepare_stream_sheet(detail, TRIAL_FIELDS, [16, 45, 18, 20, 18, 18, 16, 18, 24, 45, 14, 18, 12, 14, 24, 20, 18, 14, 14, 20])
    for row in trials:
        append_values(detail, [row.get(f) for f in TRIAL_FIELDS])

    countries = wb.create_sheet("Country_Counts")
    country_fields = ["country", "unique_study_count", "percentage_total", "percentage_known_locations"]
    prepare_stream_sheet(countries, country_fields, [30, 22, 22, 30])
    for row in country_rows:
        append_values(countries, [row.get(f) for f in country_fields], (3, 4))

    location_sheet = None
    location_sheet_number = 0
    rows_in_sheet = 0
    for row in locations:
        if location_sheet is None or rows_in_sheet >= location_rows_per_sheet:
            location_sheet_number += 1
            name = "Locations" if location_sheet_number == 1 else f"Locations_{location_sheet_number:03d}"
            location_sheet = wb.create_sheet(name)
            prepare_stream_sheet(location_sheet, LOCATION_FIELDS, [16, 16, 45, 24, 24, 18, 28])
            rows_in_sheet = 0
        append_values(location_sheet, [row.get(f) for f in LOCATION_FIELDS])
        rows_in_sheet += 1
    if location_sheet is None:
        location_sheet = wb.create_sheet("Locations")
        prepare_stream_sheet(location_sheet, LOCATION_FIELDS, [16, 16, 45, 24, 24, 18, 28])

    add_intervention_type_audit_sheet(wb, summary["intervention_type_audit"])

    defs = wb.create_sheet("Definitions")
    prepare_stream_sheet(defs, ["Item", "Definition"], [28, 100])
    time_info = summary["time_analysis"]
    for row in [["Classification model", "Version 2 — five mutually exclusive categories"], ["Unit of analysis", "One unique NCT ID"], ["Study type filter", summary["study_type_filter"]], ["Study product filter", summary["study_product_filter"]], ["Drug candidate definition", summary["drug_candidate_definition"]], ["Final intervention exclusion rule", summary["final_intervention_exclusion_rule"]], ["Temporal field", "ClinicalTrials.gov registered Start Date; not actual participant enrollment date"], ["Reference date", time_info["reference_date"]], ["RECENT_3Y boundaries", f"{time_info['boundaries']['RECENT_3Y']['start_exclusive']} < Start Date <= {time_info['boundaries']['RECENT_3Y']['end_inclusive']}"], ["YEARS_4_TO_6_AGO boundaries", f"{time_info['boundaries']['YEARS_4_TO_6_AGO']['start_exclusive']} < Start Date <= {time_info['boundaries']['YEARS_4_TO_6_AGO']['end_inclusive']}"], ["Partial dates", time_info["partial_date_rule"]], ["TIME_UNKNOWN", "Missing or unparseable Start Date"], ["TIME_AMBIGUOUS", "Partial Start Date interval crosses a time boundary"], ["Time denominator", "All unique NCT IDs whose registered Start Date is assigned unambiguously to the named cohort after the study-type and product filters"], ["UNKNOWN", "No usable registered location-country value"], ["PERMISSIBLE", "At least one usable country and no configured US country value"], ["US_ONLY", "The complete unique country set is exactly the configured US set: United States only"], ["US_NON_CHINA_MULTI", "Contains United States, does not contain China, and contains at least one additional non-US country"], ["NEXUS", "Contains at least one United States location and at least one China location; other countries may also be present"], ["NO_US", "High-level group equal to PERMISSIBLE for known-location studies"], ["HAS_US", "High-level group equal to US_ONLY + US_NON_CHINA_MULTI + NEXUS"], ["US values", ", ".join(cfg["us_countries"])], ["China values", ", ".join(cfg["china_countries"])], ["Location interpretation", "Registered or planned study facilities, not confirmed participant nationality or actual country-level enrollment."]]:
        append_values(defs, row)

    meta = wb.create_sheet("Run_Metadata"); prepare_stream_sheet(meta, ["Key", "Value"], [35, 100])
    for key, value in manifest.items():
        append_values(meta, [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for key in ("run_directory", "analysis_timestamp_utc", "reference_date", "raw_hash_verification", "summary_only", "time_analysis_status", "study_type_filter", "study_product_filter", "drug_candidate_definition", "final_intervention_exclusion_rule", "snapshot_unique_studies"):
        append_values(meta, [key, summary[key]])
    wb.save(path)


def make_summary_workbook(path: Path, summary: dict, cfg: dict, manifest: dict) -> None:
    """Create the lightweight aggregate-only workbook used by streaming mode."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Executive_Summary")
    prepare_stream_sheet(ws, ["Nexus & Permissible Summary-Only Analysis", "Value"], [62, 100])
    append_values(ws, ["Analysis mode", summary["analysis_mode"]])
    append_values(ws, ["Result status", summary["result_status"]])
    append_values(ws, ["Study type filter", summary["study_type_filter"]])
    append_values(ws, ["Study product filter", summary["study_product_filter"]])
    append_values(ws, ["Drug candidate definition", summary["drug_candidate_definition"]])
    append_values(ws, ["Analysis denominator", summary["total_studies"]])
    append_values(ws, ["Known-location denominator", summary["known_location_studies"]])
    for bucket in ("NEXUS", "PERMISSIBLE"):
        append_values(ws, [f"All-time {bucket} count", summary["buckets"][bucket]["count"]])
        append_values(ws, [f"All-time {bucket} %", summary["buckets"][bucket]["percentage_total"]], (2,))
    append_values(ws, ["Reference date", summary["reference_date"]])
    append_values(ws, ["Time analysis status", summary["time_analysis_status"]])
    for cohort, label in (("RECENT_3Y", "Recent 3 years"), ("YEARS_4_TO_6_AGO", "4–6 years ago")):
        item = summary["time_analysis"]["cohorts"][cohort]
        append_values(ws, [f"{label} denominator", item["denominator"]])
        for bucket in ("NEXUS", "PERMISSIBLE"):
            append_values(ws, [f"{label} {bucket} count", item["buckets"][bucket]["count"]])
            append_values(ws, [f"{label} {bucket} % of cohort", item["buckets"][bucket]["percentage_cohort"]], (2,))
    for row in summary["time_analysis"]["comparison"]:
        if row["metric"] in {"NEXUS % of all cohort", "PERMISSIBLE % of all cohort"}:
            append_values(ws, [f"{row['metric']} percentage-point change", row["change"]], (2,))
    if summary["all_vs_drug_diagnostic"] is not None:
        append_values(ws, ["ALL VS DRUG DIAGNOSTIC", None])
        for row in summary["all_vs_drug_diagnostic"]:
            append_values(ws, [f"{row['metric']} — All Studies", row["all_studies"]], (2,) if row["value_type"] == "percentage" else ())
            append_values(ws, [f"{row['metric']} — Industry Drug-containing Candidates", row["drug_studies"]], (2,) if row["value_type"] == "percentage" else ())
    append_values(ws, ["Location interpretation", "Registered or planned study facilities, not confirmed participant nationality or actual country-level enrollment."])

    classification = wb.create_sheet("Classification_Summary")
    prepare_stream_sheet(classification, ["bucket", "count", "percentage_total", "percentage_known_locations", "percentage_has_us", "denominator notes"], [25, 16, 22, 30, 22, 55])
    for bucket in BUCKETS:
        item = summary["buckets"][bucket]
        append_values(classification, [bucket, item["count"], item["percentage_total"], item["percentage_known_locations"], item["percentage_has_us"], "Primary denominator: unique studies after applying study_product_filter"], (3, 4, 5))

    comparison = wb.create_sheet("Time_Comparison")
    prepare_stream_sheet(comparison, ["Metric", "Recent 3 years", "4–6 years ago", "Change", "Change unit / denominator"], [38, 22, 22, 22, 55])
    for row in summary["time_analysis"]["comparison"]:
        is_percentage = row["change_unit"] == "percentage_points"
        append_values(comparison, [row["metric"], row["recent_3y"], row["years_4_to_6_ago"], row["change"], "percentage points; Recent minus 4–6 years" if is_percentage else "count; Recent minus 4–6 years"], (2, 3, 4) if is_percentage else ())

    pivot = wb.create_sheet("Time_Location_Pivot")
    pivot_buckets = ["NEXUS", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "UNKNOWN"]
    prepare_stream_sheet(pivot, ["Time cohort", "Measure", *pivot_buckets, "Total"], [25, 18, 16, 18, 16, 28, 16, 16])
    for cohort in TIME_COHORTS:
        counts = summary["time_analysis"]["time_location_counts"][cohort]
        denominator = summary["time_analysis"]["cohort_counts_all_time"][cohort]
        append_values(pivot, [cohort, "Count", *[counts[b] for b in pivot_buckets], denominator])
        if cohort in COMPARISON_COHORTS:
            item = summary["time_analysis"]["cohorts"][cohort]
            total_percentage = 1.0 if denominator else None
            append_values(pivot, [cohort, "% of cohort", *[item["buckets"][b]["percentage_cohort"] for b in pivot_buckets], total_percentage], tuple(range(3, 9)))

    add_intervention_type_audit_sheet(wb, summary["intervention_type_audit"])

    definitions = wb.create_sheet("Definitions")
    prepare_stream_sheet(definitions, ["Item", "Definition"], [34, 105])
    time_info = summary["time_analysis"]
    rows = [
        ["Start Date field path", "protocolSection.statusModule.startDateStruct.date"],
        ["Start Date type path", "protocolSection.statusModule.startDateStruct.type"],
        ["Start Date limitation", "Registered Start Date is not confirmed actual first-participant enrollment date"],
        ["Study type filter", summary["study_type_filter"]],
        ["Study product filter", summary["study_product_filter"]],
        ["Drug candidate definition", summary["drug_candidate_definition"]],
        ["Final intervention exclusion rule", summary["final_intervention_exclusion_rule"]],
        ["Reference date", summary["reference_date"]],
        ["RECENT_3Y boundaries", f"{time_info['boundaries']['RECENT_3Y']['start_exclusive']} < Start Date <= {time_info['boundaries']['RECENT_3Y']['end_inclusive']}"],
        ["YEARS_4_TO_6_AGO boundaries", f"{time_info['boundaries']['YEARS_4_TO_6_AGO']['start_exclusive']} < Start Date <= {time_info['boundaries']['YEARS_4_TO_6_AGO']['end_inclusive']}"],
        ["Partial dates", time_info["partial_date_rule"]],
        ["Time denominator", "Unique eligible NCT IDs assigned unambiguously to the named time cohort after the study-type and product filters"],
        ["UNKNOWN", "No usable registered location-country value"],
        ["PERMISSIBLE", "Known country set with no exact United States"],
        ["US_ONLY", "Complete unique country set is exactly {United States}"],
        ["US_NON_CHINA_MULTI", "Contains United States, not China, and another non-US country"],
        ["NEXUS", "Contains exact United States and exact China; other countries may be present"],
        ["US values", ", ".join(cfg["us_countries"])],
        ["China values", ", ".join(cfg["china_countries"])],
        ["Location interpretation", "Registered or planned facilities, not participant nationality or confirmed country-level enrollment"],
    ]
    for row in rows:
        append_values(definitions, row)

    metadata = wb.create_sheet("Run_Metadata")
    prepare_stream_sheet(metadata, ["Key", "Value"], [38, 105])
    values = {
        "run_directory": summary["run_directory"],
        "harvest_timestamp_utc": summary["harvest_timestamp_utc"],
        "analysis_timestamp_utc": summary["analysis_timestamp_utc"],
        "reference_date": summary["reference_date"],
        "page_count": manifest["page_count"],
        "raw_study_count": manifest["raw_study_count"],
        "unique_nct_count": manifest["unique_nct_count"],
        "analysis_denominator": summary["total_studies"],
        "raw_hash_verification": summary["raw_hash_verification"],
        "summary_only": True,
        "start_date_field_available_in_snapshot": summary["start_date_field_available_in_snapshot"],
        "study_type_filter": summary["study_type_filter"],
        "study_product_filter": summary["study_product_filter"],
        "drug_candidate_definition": summary["drug_candidate_definition"],
        "final_intervention_exclusion_rule": summary["final_intervention_exclusion_rule"],
    }
    for key, value in values.items():
        append_values(metadata, [key, value])
    wb.save(path)


def verify_summary_workbook(path: Path, summary: dict) -> None:
    def same_values(actual_values, expected_values) -> bool:
        if len(actual_values) != len(expected_values):
            return False
        for actual, expected_value in zip(actual_values, expected_values):
            if isinstance(actual, (int, float)) and isinstance(expected_value, (int, float)):
                if not math.isclose(actual, expected_value, rel_tol=1e-12, abs_tol=1e-15):
                    return False
            elif actual != expected_value:
                return False
        return True

    wb = load_workbook(path, data_only=False, read_only=True)
    expected = ["Executive_Summary", "Classification_Summary", "Time_Comparison", "Time_Location_Pivot"]
    if summary["intervention_type_audit"] is not None:
        expected.append("Intervention_Type_Audit")
    expected += ["Definitions", "Run_Metadata"]
    if wb.sheetnames != expected:
        raise AssertionError(f"Unexpected summary-only workbook sheets: {wb.sheetnames}")
    executive = {row[0].value: row[1].value for row in wb["Executive_Summary"].iter_rows(min_row=2) if len(row) >= 2}
    if executive.get("Analysis denominator") != summary["total_studies"]:
        raise AssertionError("Summary Excel total does not match summary.json")
    if executive.get("Known-location denominator") != summary["known_location_studies"]:
        raise AssertionError("Summary Excel known-location count does not match summary.json")
    if executive.get("Reference date") != summary["reference_date"]:
        raise AssertionError("Summary Excel reference date does not match summary.json")
    if executive.get("Study type filter") != summary["study_type_filter"]:
        raise AssertionError("Summary Excel study type filter does not match summary.json")
    if executive.get("Study product filter") != summary["study_product_filter"]:
        raise AssertionError("Summary Excel product filter does not match summary.json")
    if summary["intervention_type_audit"] is not None:
        audit_rows = list(wb["Intervention_Type_Audit"].iter_rows(min_row=2, values_only=True))
        exact_rows = {row[0]: row[1:] for row in audit_rows if row and row[0] not in {
            "Lead Sponsor filter", "Total Industry studies", "Industry Drug-containing studies",
            "Industry studies not containing DRUG", "Industry studies with missing intervention type",
            "Device-involved Industry Drug studies", "EXACT OBSERVED COMBINATIONS", "HIGHLIGHTED EXACT COMBINATIONS",
        }}
        for row in summary["intervention_type_audit"]["combination_counts"]:
            actual = exact_rows.get(row["combination"])
            if actual is None or not same_values(actual[:2], (row["unique_study_count"], row["percentage_industry_drug_containing"])):
                raise AssertionError(f"Summary Excel intervention combination mismatch: {row['combination']}")

    classification = {row[0]: row[1:] for row in wb["Classification_Summary"].iter_rows(min_row=2, values_only=True)}
    for bucket in BUCKETS:
        expected_values = summary["buckets"][bucket]
        actual = classification.get(bucket)
        if actual is None or not same_values(actual[:4], (
            expected_values["count"], expected_values["percentage_total"],
            expected_values["percentage_known_locations"], expected_values["percentage_has_us"],
        )):
            raise AssertionError(f"Summary Excel {bucket} aggregate mismatch")

    comparison_rows = list(wb["Time_Comparison"].iter_rows(min_row=2, values_only=True))
    if len(comparison_rows) != len(summary["time_analysis"]["comparison"]):
        raise AssertionError("Summary Excel time comparison row count mismatch")
    for actual, expected_row in zip(comparison_rows, summary["time_analysis"]["comparison"]):
        if not same_values(actual[:4], (
            expected_row["metric"], expected_row["recent_3y"],
            expected_row["years_4_to_6_ago"], expected_row["change"],
        )):
            raise AssertionError(f"Summary Excel time comparison mismatch: {expected_row['metric']}")

    pivot_rows = list(wb["Time_Location_Pivot"].iter_rows(min_row=2, values_only=True))
    if len(pivot_rows) != len(TIME_COHORTS) + len(COMPARISON_COHORTS):
        raise AssertionError("Summary Excel time pivot row count mismatch")
    pivot = {(row[0], row[1]): row[2:] for row in pivot_rows}
    for cohort in TIME_COHORTS:
        counts = summary["time_analysis"]["time_location_counts"][cohort]
        denominator = summary["time_analysis"]["cohort_counts_all_time"][cohort]
        expected_counts = tuple(counts[bucket] for bucket in BUCKETS) + (denominator,)
        if not same_values(pivot.get((cohort, "Count"), ()), expected_counts):
            raise AssertionError(f"Summary Excel {cohort} counts do not match summary.json")
        if cohort in COMPARISON_COHORTS:
            cohort_summary = summary["time_analysis"]["cohorts"][cohort]
            expected_percentages = tuple(
                cohort_summary["buckets"][bucket]["percentage_cohort"] for bucket in BUCKETS
            ) + ((1.0 if denominator else None),)
            if not same_values(pivot.get((cohort, "% of cohort"), ()), expected_percentages):
                raise AssertionError(f"Summary Excel {cohort} percentages do not match summary.json")

    definitions = {row[0].value: row[1].value for row in wb["Definitions"].iter_rows(min_row=2)}
    if definitions.get("Reference date") != summary["reference_date"]:
        raise AssertionError("Summary Excel Definitions reference date mismatch")
    metadata = {row[0].value: row[1].value for row in wb["Run_Metadata"].iter_rows(min_row=2)}
    expected_metadata = {
        "run_directory": summary["run_directory"],
        "harvest_timestamp_utc": summary["harvest_timestamp_utc"],
        "analysis_timestamp_utc": summary["analysis_timestamp_utc"],
        "reference_date": summary["reference_date"],
        "unique_nct_count": summary["snapshot_unique_studies"],
        "analysis_denominator": summary["total_studies"],
        "raw_hash_verification": summary["raw_hash_verification"],
        "summary_only": True,
        "study_type_filter": summary["study_type_filter"],
        "study_product_filter": summary["study_product_filter"],
        "drug_candidate_definition": summary["drug_candidate_definition"],
        "final_intervention_exclusion_rule": summary["final_intervention_exclusion_rule"],
    }
    for key, expected_value in expected_metadata.items():
        if metadata.get(key) != expected_value:
            raise AssertionError(f"Summary Excel metadata mismatch: {key}")
    if any(name in wb.sheetnames for name in ("Study_Detail", "Locations", "Country_Counts")):
        raise AssertionError("Summary-only workbook contains a prohibited detail sheet")
    wb.close()


def verify_summary_text_outputs(out: Path, summary: dict) -> None:
    """Verify that JSON and Markdown expose the same aggregate values used by Excel."""
    written = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    if written != summary:
        raise AssertionError("Written summary.json does not match the in-memory summary")
    markdown = (out / "summary.md").read_text(encoding="utf-8")
    required_fragments = [
        f"Analysis mode: **{summary['analysis_mode']}**",
        f"Study type filter: **{summary['study_type_filter']}**",
        f"Study product filter: **{summary['study_product_filter']}**",
        f"Drug candidate definition: **{summary['drug_candidate_definition']}**",
        f"Total studies: **{summary['total_studies']:,}**",
        f"Known-location studies: **{summary['known_location_studies']:,}**",
        f"Reference date: **{summary['reference_date']}**" if summary["time_analysis_status"] == "AVAILABLE" else "**UNAVAILABLE:**",
    ]
    for bucket in BUCKETS:
        required_fragments.append(f"{bucket}: **{summary['buckets'][bucket]['count']:,}")
    audit = summary["intervention_type_audit"]
    if audit is not None:
        required_fragments.extend([
            f"Total Industry studies: **{audit['total_industry_studies']:,}**",
            f"Industry Drug-containing studies: **{audit['industry_drug_containing_studies']:,}**",
            f"Device-involved Industry Drug studies: **{audit['device_involved_drug_studies']:,}**",
        ])
        for row in audit["combination_counts"]:
            required_fragments.append(f"| {row['combination']} | {row['unique_study_count']:,} |")
    if summary["time_analysis_status"] == "AVAILABLE":
        fmt = lambda value: "N/A" if value is None else f"{value:.2%}"
        for cohort in TIME_COHORTS:
            required_fragments.append(
                f"| {cohort} | {summary['time_analysis']['cohort_counts_all_time'][cohort]:,} |"
            )
        for cohort in COMPARISON_COHORTS:
            item = summary["time_analysis"]["cohorts"][cohort]
            for bucket in BUCKETS:
                bucket_item = item["buckets"][bucket]
                required_fragments.append(
                    f"| {cohort} | {bucket} | {bucket_item['count']:,} | "
                    f"{fmt(bucket_item['percentage_cohort'])} | {fmt(bucket_item['percentage_known_locations'])} |"
                )
        for comparison_row in summary["time_analysis"]["comparison"]:
            is_percentage = comparison_row["change_unit"] == "percentage_points"
            recent_value = fmt(comparison_row["recent_3y"]) if is_percentage else f"{comparison_row['recent_3y']:,}"
            older_value = fmt(comparison_row["years_4_to_6_ago"]) if is_percentage else f"{comparison_row['years_4_to_6_ago']:,}"
            change_value = fmt(comparison_row["change"]) if is_percentage else f"{comparison_row['change']:+,}"
            required_fragments.append(
                f"| {comparison_row['metric']} | {recent_value} | {older_value} | {change_value} | "
                f"{'percentage points' if is_percentage else 'count'} |"
            )
    missing = [fragment for fragment in required_fragments if fragment not in markdown]
    if missing:
        raise AssertionError(f"summary.md is missing aggregate values: {missing}")


def verify_workbook(path: Path, summary: dict):
    wb = load_workbook(path, data_only=False, read_only=True)
    fixed = ["Executive_Summary", "Classification_Summary", "Time_Comparison", "Time_Location_Pivot", "Study_Detail", "Country_Counts"]
    location_sheets = [name for name in wb.sheetnames if name == "Locations" or name.startswith("Locations_")]
    required_tail = ["Definitions", "Run_Metadata"]
    if wb.sheetnames[:6] != fixed or not location_sheets or wb.sheetnames[-2:] != required_tail:
        raise AssertionError(f"Unexpected workbook sheets: {wb.sheetnames}")
    values = {row[0].value: row[1].value for row in wb["Executive_Summary"].iter_rows(min_row=2) if len(row) >= 2}
    if values["Total unique studies"] != summary["total_studies"]:
        raise AssertionError("Excel total does not match summary.json")
    classification = {row[0].value: row[1].value for row in wb["Classification_Summary"].iter_rows(min_row=2)}
    for bucket in BUCKETS:
        if classification.get(bucket) != summary["buckets"][bucket]["count"]:
            raise AssertionError(f"Excel {bucket} count does not match summary.json")
    if values["HAS_US count"] != summary["us_presence_groups"]["HAS_US"]["count"]:
        raise AssertionError("Excel HAS_US count does not match summary.json")
    time_comparison = list(wb["Time_Comparison"].iter_rows(min_row=2, values_only=True))
    if len(time_comparison) != len(summary["time_analysis"]["comparison"]):
        raise AssertionError("Excel time comparison row count does not match summary.json")
    pivot = list(wb["Time_Location_Pivot"].iter_rows(min_row=2, values_only=True))
    if len(pivot) != len(TIME_COHORTS) + len(COMPARISON_COHORTS):
        raise AssertionError("Excel time pivot must contain all cohort counts plus primary-cohort percentages")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("config.json"))
    parser.add_argument("--output-name", default="out_v2", help="Versioned output directory name under the selected run")
    parser.add_argument("--reference-date", help="ISO date; defaults to manifest harvest date (UTC)")
    parser.add_argument("--summary-only", action="store_true", help="Stream raw pages and write aggregate-only JSON, Markdown, and Excel")
    parser.add_argument("--study-type", choices=("all", "interventional"), default="all",
                        help="Restrict the denominator to StudyType == INTERVENTIONAL")
    parser.add_argument("--study-product", choices=("all", "drug"), default="all",
                        help="Restrict the denominator to Industry-sponsored studies containing a DRUG intervention")
    args = parser.parse_args(argv)
    cfg = json.loads(args.config.read_text(encoding="utf-8"))
    manifest_path = args.run / "manifest.json"
    if not manifest_path.exists():
        parser.error("selected run has no manifest.json (counts-only or failed runs are rejected)")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not manifest.get("files") or not isinstance(manifest.get("snapshot_complete"), bool):
        parser.error("manifest is not a verifiable raw snapshot")
    if len(manifest["files"]) != manifest.get("page_count"):
        parser.error("manifest page_count disagrees with files list")
    if manifest.get("snapshot_complete") and manifest.get("next_page_token_remaining") is not False:
        parser.error("complete snapshot manifest must explicitly confirm that no next-page token remains")
    if args.summary_only and not manifest["snapshot_complete"]:
        parser.error("--summary-only requires a complete snapshot; partial smoke snapshots are rejected")
    requested_fields = {
        field.strip() for field in manifest.get("request_parameters", {}).get("fields", "").split(",") if field.strip()
    }
    candidate_fields_available = {"InterventionType", "LeadSponsorClass"}.issubset(requested_fields)
    if args.study_product == "drug" and not candidate_fields_available:
        parser.error("--study-product drug requires InterventionType and LeadSponsorClass in the snapshot manifest; harvest a new snapshot with those structured fields")
    try:
        reference_date = (datetime.strptime(args.reference_date, "%Y-%m-%d").date() if args.reference_date
                          else datetime.fromisoformat(manifest["harvest_timestamp_utc"].replace("Z", "+00:00")).date())
    except (KeyError, ValueError) as exc:
        parser.error(f"invalid reference date or manifest harvest timestamp: {exc}")
    out = args.run / args.output_name
    if out.exists() and any(out.iterdir()):
        parser.error(f"output directory already exists and is not empty: {out}")

    analysis_timestamp = datetime.now(timezone.utc).isoformat()
    duplicate_ids = Counter()
    raw_count = 0

    if args.summary_only:
        seen_nct_ids = set()
        counts = Counter()
        time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
        all_time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
        drug_time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
        intervention_audit_counts = Counter()
        intervention_combinations = Counter()
        study_type_total = 0
        for entry in manifest["files"]:
            path = args.run / entry["path"]
            if not path.is_file() or sha256(path) != entry["sha256"]:
                raise RuntimeError(f"Raw-file hash verification failed: {path}")
            payload = json.loads(path.read_text(encoding="utf-8"))
            studies = payload.get("studies")
            if not isinstance(studies, list) or len(studies) != entry["study_count"]:
                raise RuntimeError(f"Raw page study count disagrees with manifest: {path}")
            raw_count += len(studies)
            for study in studies:
                nct = study.get("protocolSection", {}).get("identificationModule", {}).get("nctId")
                if not nct:
                    raise ValueError("Every retained record must have an NCT ID")
                if nct in seen_nct_ids:
                    duplicate_ids[nct] += 1
                    continue
                seen_nct_ids.add(nct)
                type_eligible = study_type_is_eligible(study, args.study_type)
                if type_eligible:
                    study_type_total += 1
                    update_intervention_type_audit(intervention_audit_counts, intervention_combinations, study)
                eligible = type_eligible and product_is_eligible(study, args.study_product)
                row, _ = extract(study, cfg, reference_date, include_locations=False)
                validate_trial_row(row, cfg)
                if type_eligible:
                    all_time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
                if type_eligible and product_is_eligible(study, "drug"):
                    drug_time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
                if eligible:
                    counts[row["bucket"]] += 1
                    time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
        snapshot_total = len(seen_nct_ids)
        if raw_count != manifest["raw_study_count"]:
            raise RuntimeError("Raw study count disagrees with manifest")
        if snapshot_total != manifest["unique_nct_count"]:
            raise RuntimeError("Unique NCT count disagrees with manifest")
        if raw_count - snapshot_total != manifest.get("duplicate_nct_count", raw_count - snapshot_total):
            raise RuntimeError("Duplicate NCT count disagrees with manifest")
        total = sum(counts.values())
        time_analysis = build_time_analysis_from_counts(time_bucket_counts, reference_date, total)
        all_time_analysis = build_time_analysis_from_counts(all_time_bucket_counts, reference_date, study_type_total)
        candidate_total = intervention_audit_counts["industry_drug_containing"]
        drug_time_analysis = (build_time_analysis_from_counts(
            drug_time_bucket_counts, reference_date, candidate_total
        ) if candidate_fields_available and candidate_total else None)
        diagnostic = build_all_vs_drug_diagnostic(all_time_analysis, drug_time_analysis) if drug_time_analysis else None
        intervention_audit = (build_intervention_type_audit(
            intervention_audit_counts, intervention_combinations, args.study_type
        ) if candidate_fields_available else None)
        summary = build_summary(counts=counts, total=total, time_analysis=time_analysis, manifest=manifest, cfg=cfg,
                                duplicate_ids=duplicate_ids, reference_date=reference_date, run_directory=args.run,
                                summary_only=True, analysis_timestamp_utc=analysis_timestamp,
                                study_product=args.study_product, study_type=args.study_type,
                                intervention_type_audit=intervention_audit,
                                all_vs_drug_diagnostic=diagnostic)
        out.mkdir(exist_ok=True)
        write_summary_files(out, summary, cfg, manifest)
        workbook_path = out / "nexus_permissible_summary.xlsx"
        make_summary_workbook(workbook_path, summary, cfg, manifest)
        verify_summary_text_outputs(out, summary)
        verify_summary_workbook(workbook_path, summary)
        prohibited = ["trials.csv", "locations_long.csv", "country_counts.csv", "country_vocabulary.csv", "nexus_permissible_results.xlsx"]
        if any((out / name).exists() for name in prohibited):
            raise AssertionError("Summary-only mode generated a prohibited detail artifact")
        print(out)
        return 0

    studies_by_id = {}
    for entry in manifest["files"]:
        path = args.run / entry["path"]
        if not path.is_file() or sha256(path) != entry["sha256"]:
            raise RuntimeError(f"Raw-file hash verification failed: {path}")
        payload = json.loads(path.read_text(encoding="utf-8"))
        studies = payload.get("studies")
        if not isinstance(studies, list) or len(studies) != entry["study_count"]:
            raise RuntimeError(f"Raw page study count disagrees with manifest: {path}")
        raw_count += len(studies)
        for study in studies:
            nct = study.get("protocolSection", {}).get("identificationModule", {}).get("nctId")
            if not nct:
                raise ValueError("Every retained record must have an NCT ID")
            if nct in studies_by_id:
                duplicate_ids[nct] += 1
            else:
                studies_by_id[nct] = study
    if raw_count != manifest["raw_study_count"]:
        raise RuntimeError("Partial-download detection failed: manifest raw count disagrees with page files")
    if len(studies_by_id) != manifest["unique_nct_count"]:
        raise RuntimeError("Unique NCT count disagrees with manifest")
    if raw_count - len(studies_by_id) != manifest.get("duplicate_nct_count", raw_count - len(studies_by_id)):
        raise RuntimeError("Duplicate NCT count disagrees with manifest")
    trials, locations = [], []
    all_time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
    drug_time_bucket_counts = {cohort: Counter() for cohort in TIME_COHORTS}
    intervention_audit_counts = Counter()
    intervention_combinations = Counter()
    study_type_total = 0
    for nct in sorted(studies_by_id):
        study = studies_by_id[nct]
        type_eligible = study_type_is_eligible(study, args.study_type)
        if type_eligible:
            study_type_total += 1
            update_intervention_type_audit(intervention_audit_counts, intervention_combinations, study)
        eligible = type_eligible and product_is_eligible(study, args.study_product)
        row, locs = extract(study, cfg, reference_date, include_locations=eligible)
        validate_trial_row(row, cfg)
        if type_eligible:
            all_time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
        if type_eligible and product_is_eligible(study, "drug"):
            drug_time_bucket_counts[row["time_cohort"]][row["bucket"]] += 1
        if eligible:
            trials.append(row)
            locations.extend(locs)
    counts = Counter(row["bucket"] for row in trials)
    total = len(trials)
    time_analysis = build_time_analysis(trials, reference_date)
    all_time_analysis = build_time_analysis_from_counts(all_time_bucket_counts, reference_date, study_type_total)
    candidate_total = intervention_audit_counts["industry_drug_containing"]
    drug_time_analysis = (build_time_analysis_from_counts(
        drug_time_bucket_counts, reference_date, candidate_total
    ) if candidate_fields_available and candidate_total else None)
    diagnostic = build_all_vs_drug_diagnostic(all_time_analysis, drug_time_analysis) if drug_time_analysis else None
    intervention_audit = (build_intervention_type_audit(
        intervention_audit_counts, intervention_combinations, args.study_type
    ) if candidate_fields_available else None)
    summary = build_summary(counts=counts, total=total, time_analysis=time_analysis, manifest=manifest, cfg=cfg,
                            duplicate_ids=duplicate_ids, reference_date=reference_date, run_directory=args.run,
                            summary_only=False, analysis_timestamp_utc=analysis_timestamp,
                            study_product=args.study_product, study_type=args.study_type,
                            intervention_type_audit=intervention_audit,
                            all_vs_drug_diagnostic=diagnostic)
    known = summary["known_location_studies"]
    countries_to_ids = defaultdict(set)
    for row in trials:
        for country in filter(None, row["countries_pipe"].split("|")):
            countries_to_ids[country].add(row["nct_id"])
    country_rows = [{"country": c, "unique_study_count": len(ids), "percentage_total": pct(len(ids), total),
                     "percentage_known_locations": pct(len(ids), known)} for c, ids in sorted(countries_to_ids.items(), key=lambda x: (-len(x[1]), x[0]))]
    out.mkdir(exist_ok=True)
    write_csv(out / "trials.csv", TRIAL_FIELDS, trials)
    write_csv(out / "locations_long.csv", LOCATION_FIELDS, locations)
    write_csv(out / "country_counts.csv", ["country", "unique_study_count", "percentage_total", "percentage_known_locations"], country_rows)
    write_csv(out / "country_vocabulary.csv", ["country"], ({"country": c} for c in sorted(countries_to_ids)))
    write_summary_files(out, summary, cfg, manifest)
    verify_summary_text_outputs(out, summary)
    workbook_path = out / "nexus_permissible_results.xlsx"
    make_workbook(workbook_path, summary, trials, locations, country_rows, cfg, manifest,
                  location_rows_per_sheet=cfg.get("excel_location_rows_per_sheet", 500_000))
    verify_workbook(workbook_path, summary)
    print(out)
    return 0


if __name__ == "__main__": raise SystemExit(main())
