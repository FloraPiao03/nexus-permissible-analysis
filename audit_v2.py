#!/usr/bin/env python3
"""Independent five-category audit; deliberately imports no production modules."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

BUCKETS = ("UNKNOWN", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "NEXUS")
US = "United States"
CHINA = "China"


def classify(countries: set[str]) -> tuple[str, str]:
    if not countries:
        return "UNKNOWN", "UNKNOWN"
    if US not in countries:
        return "PERMISSIBLE", "NO_US"
    if countries == {US}:
        return "US_ONLY", "HAS_US"
    if CHINA in countries:
        return "NEXUS", "HAS_US"
    return "US_NON_CHINA_MULTI", "HAS_US"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def excel_values_match(expected, actual) -> bool:
    if actual is None or len(expected) != len(actual):
        return False
    for left, right in zip(expected, actual):
        if left is None or right is None:
            if left is not right:
                return False
        elif isinstance(left, float) or isinstance(right, float):
            if abs(left - right) > 1e-15:
                return False
        elif left != right:
            return False
    return True


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", type=Path, required=True)
    parser.add_argument("--new-output", default="out_v2")
    parser.add_argument("--old-output", default="out")
    args = parser.parse_args()
    run = args.run
    new_out = run / args.new_output
    old_out = run / args.old_output

    manifest = json.loads((run / "manifest.json").read_text(encoding="utf-8"))
    hash_errors = []
    seen = set()
    expected = {}
    raw_bucket_counts = Counter()
    group_counts = Counter()
    country_location_rows = Counter()
    country_nct_ids = defaultdict(set)
    whitespace_rows = Counter()
    case_variants = defaultdict(Counter)

    for entry in manifest["files"]:
        page = run / entry["path"]
        if sha256(page) != entry["sha256"]:
            hash_errors.append(str(page))
        payload = json.loads(page.read_text(encoding="utf-8"))
        for study in payload.get("studies", []):
            protocol = study.get("protocolSection", {})
            nct = protocol.get("identificationModule", {}).get("nctId")
            if not nct or nct in seen:
                continue
            seen.add(nct)
            normalized = set()
            for location in protocol.get("contactsLocationsModule", {}).get("locations") or []:
                value = location.get("country")
                if not isinstance(value, str):
                    continue
                country_location_rows[value] += 1
                country_nct_ids[value].add(nct)
                if value != value.strip():
                    whitespace_rows[value] += 1
                if value.strip():
                    case_variants[value.strip().casefold()][value] += 1
                    normalized.add(value.strip())
            bucket, group = classify(normalized)
            expected[nct] = (bucket, group)
            raw_bucket_counts[bucket] += 1
            group_counts[group] += 1

    actual = {}
    with (new_out / "trials.csv").open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            actual[row["nct_id"]] = (row["bucket"], row["us_presence_group"])
    missing_ids = sorted(set(expected) - set(actual))
    extra_ids = sorted(set(actual) - set(expected))
    row_mismatches = [(nct, expected[nct], actual[nct]) for nct in set(expected) & set(actual) if expected[nct] != actual[nct]]

    old_buckets = {}
    with (old_out / "trials.csv").open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            old_buckets[row["nct_id"]] = row["bucket"]
    transitions = Counter((old_buckets[nct], expected[nct][0]) for nct in expected)

    summary = json.loads((new_out / "summary.json").read_text(encoding="utf-8"))
    summary_errors = []
    total = len(expected)
    known = total - raw_bucket_counts["UNKNOWN"]
    has_us = group_counts["HAS_US"]
    if summary["total_studies"] != len(expected):
        summary_errors.append(["total_studies", len(expected), summary["total_studies"]])
    for bucket in BUCKETS:
        item = summary["buckets"][bucket]
        expected_values = {
            "count": raw_bucket_counts[bucket],
            "percentage_total": raw_bucket_counts[bucket] / total,
            "percentage_known_locations": None if bucket == "UNKNOWN" else raw_bucket_counts[bucket] / known,
            "percentage_has_us": raw_bucket_counts[bucket] / has_us if bucket in {"US_ONLY", "US_NON_CHINA_MULTI", "NEXUS"} else None,
        }
        for metric, expected_value in expected_values.items():
            if item[metric] != expected_value:
                summary_errors.append([f"{bucket}.{metric}", expected_value, item[metric]])
    for group in ("UNKNOWN", "NO_US", "HAS_US"):
        if summary["us_presence_groups"][group]["count"] != group_counts[group]:
            summary_errors.append([group, group_counts[group], summary["us_presence_groups"][group]["count"]])

    workbook = load_workbook(new_out / "nexus_permissible_results.xlsx", read_only=True, data_only=False)
    excel_rows = {row[0].value: tuple(cell.value for cell in row[1:5]) for row in workbook["Classification_Summary"].iter_rows(min_row=2) if len(row) >= 5}
    excel_errors = []
    for bucket in BUCKETS:
        expected_excel = (raw_bucket_counts[bucket], raw_bucket_counts[bucket] / total,
                          None if bucket == "UNKNOWN" else raw_bucket_counts[bucket] / known,
                          raw_bucket_counts[bucket] / has_us if bucket in {"US_ONLY", "US_NON_CHINA_MULTI", "NEXUS"} else None)
        if not excel_values_match(expected_excel, excel_rows.get(bucket)):
            excel_errors.append([bucket, expected_excel, excel_rows.get(bucket)])
    executive = {row[0].value: row[1].value for row in workbook["Executive_Summary"].iter_rows(min_row=2) if len(row) >= 2}
    for label, expected_value in (("Total unique studies", total), ("Known-location studies", known), ("Unknown-location studies", raw_bucket_counts["UNKNOWN"]), ("NO_US / PERMISSIBLE count", group_counts["NO_US"]), ("HAS_US count", has_us)):
        if executive.get(label) != expected_value:
            excel_errors.append([label, expected_value, executive.get(label)])
    summary_md = (new_out / "summary.md").read_text(encoding="utf-8")
    markdown_errors = [bucket for bucket in BUCKETS if f"{raw_bucket_counts[bucket]:,}" not in summary_md]

    requested_aliases = [
        "United States", "US", "U.S.", "USA", "U.S.A.", "United States of America", "the US", "America",
        "China", "PRC", "P.R.C.", "People's Republic of China", "People’s Republic of China", "Mainland China",
    ]
    plausible_pattern = re.compile(r"(^|\b)(u\.?s\.?a?\.?|united states|america|china|p\.?r\.?c\.?|people.?s republic|mainland)(\b|$)", re.I)
    plausible = sorted(value for value in country_location_rows if plausible_pattern.search(value))
    audit_rows = []
    for value in sorted(set(requested_aliases) | set(plausible)):
        relation = "canonical_us" if value == US else "canonical_china" if value == CHINA else "distinct_jurisdiction" if value == "United States Minor Outlying Islands" else "candidate_alias_not_observed" if value not in country_location_rows else "observed_candidate_review_required"
        audit_rows.append({"exact_country_string": value, "observed": value in country_location_rows,
                           "unique_nct_count": len(country_nct_ids[value]), "location_row_count": country_location_rows[value],
                           "trimmed_value": value.strip(), "has_leading_or_trailing_whitespace": value != value.strip(),
                           "relationship": relation})
    with (new_out / "country_string_audit.csv").open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(audit_rows[0]))
        writer.writeheader(); writer.writerows(audit_rows)

    report = {
        "classification_model_version": 2,
        "independence": "does not import harvest.py or analyze.py",
        "run_directory": str(run.resolve()),
        "snapshot_complete": manifest.get("snapshot_complete"),
        "next_page_token_remaining": manifest.get("next_page_token_remaining"),
        "page_count": len(manifest["files"]), "hash_error_count": len(hash_errors),
        "total_unique_nct_ids": total, "known_location_studies": known,
        "independent_bucket_counts": dict(raw_bucket_counts), "independent_us_presence_groups": dict(group_counts),
        "reconciliation": {
            "five_bucket_sum": sum(raw_bucket_counts[b] for b in BUCKETS),
            "permissible_plus_has_us": raw_bucket_counts["PERMISSIBLE"] + has_us,
            "us_subcategories_sum": raw_bucket_counts["US_ONLY"] + raw_bucket_counts["US_NON_CHINA_MULTI"] + raw_bucket_counts["NEXUS"],
        },
        "percentages": {
            bucket: {"all_studies": raw_bucket_counts[bucket] / total,
                     "known_locations": None if bucket == "UNKNOWN" else raw_bucket_counts[bucket] / known,
                     "has_us": raw_bucket_counts[bucket] / has_us if bucket in {"US_ONLY", "US_NON_CHINA_MULTI", "NEXUS"} else None}
            for bucket in BUCKETS
        },
        "old_to_new_transitions": [{"old_bucket": old, "new_bucket": new, "count": count} for (old, new), count in sorted(transitions.items())],
        "country_string_audit": {
            "observed_plausible_us_or_china_strings": plausible,
            "us_case_or_punctuation_variants": dict(case_variants[US.casefold()]),
            "china_case_or_punctuation_variants": dict(case_variants[CHINA.casefold()]),
            "whitespace_variant_location_rows": dict(whitespace_rows),
            "exact_matching_sufficient_for_snapshot": plausible == ["China", "United States", "United States Minor Outlying Islands"] and set(case_variants[US.casefold()]) == {US} and set(case_variants[CHINA.casefold()]) == {CHINA},
        },
        "comparison": {"missing_trial_ids": missing_ids, "extra_trial_ids": extra_ids,
                       "trial_row_mismatch_count": len(row_mismatches), "trial_row_mismatches_first_100": row_mismatches[:100],
                       "summary_json_errors": summary_errors, "summary_md_errors": markdown_errors, "excel_errors": excel_errors},
        "verdict": "PASS" if not (hash_errors or missing_ids or extra_ids or row_mismatches or summary_errors or markdown_errors or excel_errors) else "FAIL",
    }
    (new_out / "reclassification_audit.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0 if report["verdict"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
