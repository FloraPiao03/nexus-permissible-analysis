#!/usr/bin/env python3
"""Independent reconciliation of Interventional-only artifacts; imports no production logic."""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

BUCKETS = ("UNKNOWN", "PERMISSIBLE", "US_ONLY", "US_NON_CHINA_MULTI", "NEXUS")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    counts = Counter()
    groups = Counter()
    missing_output_ids = []
    extra_output_ids = []
    row_mismatches = []
    total = 0
    with args.source.open(newline="", encoding="utf-8") as source_fh, (args.outdir / "trials_interventional.csv").open(newline="", encoding="utf-8") as output_fh:
        source_reader = (row for row in csv.DictReader(source_fh) if row["study_type"] == "INTERVENTIONAL")
        output_reader = csv.DictReader(output_fh)
        while True:
            source_row = next(source_reader, None)
            output_row = next(output_reader, None)
            if source_row is None and output_row is None:
                break
            if source_row is None:
                extra_output_ids.append(output_row["nct_id"])
                continue
            total += 1
            counts[source_row["bucket"]] += 1
            groups[source_row["us_presence_group"]] += 1
            if output_row is None:
                missing_output_ids.append(source_row["nct_id"])
            elif source_row["nct_id"] != output_row["nct_id"]:
                row_mismatches.append({"source": source_row["nct_id"], "output": output_row["nct_id"]})
            elif source_row != output_row:
                row_mismatches.append({"nct_id": source_row["nct_id"], "reason": "field mismatch"})
    known = total - counts["UNKNOWN"]
    has_us = counts["US_ONLY"] + counts["US_NON_CHINA_MULTI"] + counts["NEXUS"]
    summary = json.loads((args.outdir / "summary_interventional.json").read_text(encoding="utf-8"))
    json_errors = []
    if summary["total_studies"] != total or summary["known_location_studies"] != known:
        json_errors.append("denominator mismatch")
    for bucket in BUCKETS:
        if summary["buckets"][bucket]["count"] != counts[bucket]:
            json_errors.append(f"{bucket} count mismatch")
    wb = load_workbook(args.outdir / "nexus_permissible_interventional.xlsx", read_only=True, data_only=False)
    excel_counts = {row[0].value: row[1].value for row in wb["Classification_Summary"].iter_rows(min_row=2)}
    excel_errors = [f"{bucket} count mismatch" for bucket in BUCKETS if excel_counts.get(bucket) != counts[bucket]]
    detail_rows = sum(1 for _ in wb["Study_Detail"].iter_rows(min_row=2, values_only=True))
    if detail_rows != total:
        excel_errors.append("Study_Detail row count mismatch")
    wb.close()
    differences = {
        "five_buckets_minus_total": sum(counts[b] for b in BUCKETS) - total,
        "permissible_plus_has_us_minus_known": counts["PERMISSIBLE"] + has_us - known,
        "us_subcategories_minus_has_us": counts["US_ONLY"] + counts["US_NON_CHINA_MULTI"] + counts["NEXUS"] - has_us,
        "group_unknown_minus_bucket_unknown": groups["UNKNOWN"] - counts["UNKNOWN"],
        "group_no_us_minus_permissible": groups["NO_US"] - counts["PERMISSIBLE"],
        "group_has_us_minus_us_subcategories": groups["HAS_US"] - has_us,
    }
    report = {
        "independence": "reads the study-level CSV directly and imports neither analyze.py nor analyze_interventional.py",
        "source": str(args.source.resolve()),
        "total_interventional": total,
        "known_location_interventional": known,
        "bucket_counts": {bucket: counts[bucket] for bucket in BUCKETS},
        "us_presence_groups": {key: groups[key] for key in ("UNKNOWN", "NO_US", "HAS_US")},
        "reconciliation_differences": differences,
        "missing_output_ids": missing_output_ids[:100],
        "extra_output_ids": extra_output_ids[:100],
        "row_mismatch_count": len(row_mismatches),
        "json_errors": json_errors,
        "excel_errors": excel_errors,
    }
    report["verdict"] = "PASS" if not any(differences.values()) and not report["missing_output_ids"] and not report["extra_output_ids"] and not row_mismatches and not json_errors and not excel_errors else "FAIL"
    target = args.outdir / "interventional_audit.json"
    target.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0 if report["verdict"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
